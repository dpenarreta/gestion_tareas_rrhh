"""Cobertura del importador/plantilla de Tareas por Excel — sub-fase 3e
(ver docs/AUDIT_LOG.md § 2026-08-07)."""

from datetime import date
from io import BytesIO

import openpyxl
import pytest
from django.contrib.auth.models import Group
from django.core.files.uploadedfile import SimpleUploadedFile
from rest_framework.test import APIClient

from apps.tasks.models import Task
from apps.users.models import User

pytestmark = pytest.mark.django_db

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER = [
    "Título",
    "Descripción",
    "Prioridad",
    "Frecuencia",
    "Fecha Inicio",
    "Fecha Fin",
    "Tiempo Objetivo",
    "Asignado a",
    "Tipo",
]


@pytest.fixture
def actor():
    # Grupo real (ver docs/AUDIT_LOG.md § 2026-09-01, NEXO-01): la
    # importación por Excel ahora valida `assigned_to` contra la jerarquía
    # visible del actor — un usuario sin grupo no puede asignarle tareas a
    # nadie más que a sí mismo.
    user = User.objects.create_user(
        username="importer", email="importer@example.com", password="Sup3r-Secr3t!"
    )
    user.groups.set([Group.objects.get(name="ANALISTA_CC")])
    return user


@pytest.fixture
def actor_client(actor):
    client = APIClient()
    client.force_authenticate(user=actor)
    return client


def _xlsx_bytes(rows: list[list]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _upload(rows: list[list]) -> SimpleUploadedFile:
    return SimpleUploadedFile("tasks.xlsx", _xlsx_bytes(rows), content_type=XLSX_CONTENT_TYPE)


def _row(**overrides) -> list:
    fields = {
        "title": "Informe",
        "description": "desc",
        "priority": "ALTA",
        "frequency": "MENSUAL",
        "start": "2026-07-01",
        "end": "2026-07-15",
        "hours": "8",
        "email": "",
        "type": "FIJA",
    }
    fields.update(overrides)
    return [
        fields["title"],
        fields["description"],
        fields["priority"],
        fields["frequency"],
        fields["start"],
        fields["end"],
        fields["hours"],
        fields["email"],
        fields["type"],
    ]


# --- plantilla -----------------------------------------------------------


def test_template_requires_auth():
    response = APIClient().get("/api/v1/tasks/template/")
    assert response.status_code == 401


def test_template_returns_expected_xlsx(actor_client):
    response = actor_client.get("/api/v1/tasks/template/")

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert "plantilla_tareas.xlsx" in response["Content-Disposition"]

    workbook = openpyxl.load_workbook(BytesIO(response.content))
    sheet = workbook["Tareas"]
    assert sheet.cell(row=1, column=1).value == "Título"
    assert sheet.cell(row=2, column=3).value == "ALTA"
    assert sheet.cell(row=2, column=9).value == "FIJA"


# --- import: autenticación y validación de archivo ------------------------


def test_import_requires_auth():
    response = APIClient().post(
        "/api/v1/tasks/import/", {"file": _upload([HEADER, _row()])}, format="multipart"
    )
    assert response.status_code == 401


def test_import_without_file_returns_400(actor_client):
    response = actor_client.post("/api/v1/tasks/import/", {}, format="multipart")
    assert response.status_code == 400


# --- import: fila válida --------------------------------------------------


def test_import_valid_row_without_email_self_assigns(actor_client, actor):
    upload = _upload([HEADER, _row()])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")

    assert response.status_code == 200
    assert response.data == {"imported": 1, "errors": []}
    task = Task.objects.get(title="Informe")
    assert task.assigned_to_id == actor.id
    assert task.created_by_id == actor.id
    assert task.status == "PENDIENTE"
    assert task.real_hours == 0
    assert task.start_date.date() == date(2026, 7, 1)
    assert task.end_date.date() == date(2026, 7, 15)


def test_import_resolves_assignee_by_email(actor_client):
    other = User.objects.create_user(
        username="other", email="other@example.com", password="Sup3r-Secr3t!"
    )
    other.groups.set([Group.objects.get(name="ANALISTA_CC")])
    upload = _upload([HEADER, _row(email="other@example.com")])

    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")

    assert response.data == {"imported": 1, "errors": []}
    assert Task.objects.get(title="Informe").assigned_to_id == other.id


def test_import_rejects_assignee_outside_visible_hierarchy(actor_client):
    """Hallazgo real de la auditoría de seguridad (ver docs/AUDIT_LOG.md §
    2026-09-01, NEXO-01) — mismo bug que la creación/edición directa de
    tareas, alcanzable también vía importación masiva por Excel."""
    outsider = User.objects.create_user(
        username="outsider", email="outsider@example.com", password="Sup3r-Secr3t!"
    )
    outsider.groups.set([Group.objects.get(name="ADMINISTRADOR")])
    upload = _upload([HEADER, _row(email="outsider@example.com")])

    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")

    assert response.data == {
        "imported": 0,
        "errors": [{"row": 2, "error": 'No podés asignar tareas a "outsider@example.com"'}],
    }
    assert not Task.objects.filter(title="Informe").exists()


def test_import_skips_fully_empty_rows(actor_client):
    upload = _upload([HEADER, [None] * 9, _row()])

    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")

    assert response.data == {"imported": 1, "errors": []}


# --- import: validaciones de fila (rechazo individual, no aborta el batch) --


def test_import_missing_title(actor_client):
    upload = _upload([HEADER, _row(title="")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data == {"imported": 0, "errors": [{"row": 2, "error": "Título requerido"}]}


def test_import_invalid_priority(actor_client):
    upload = _upload([HEADER, _row(priority="URGENTE")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [
        {"row": 2, "error": 'Prioridad inválida: "URGENTE". Use ALTA, MEDIA o BAJA'}
    ]


def test_import_invalid_frequency(actor_client):
    upload = _upload([HEADER, _row(frequency="URGENTE")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [
        {
            "row": 2,
            "error": 'Frecuencia inválida: "URGENTE". Use MENSUAL, SEMANAL, DIARIA, QUINCENAL o PUNTUAL',
        }
    ]


def test_import_missing_start_date(actor_client):
    upload = _upload([HEADER, _row(start="")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [{"row": 2, "error": "la fecha de inicio es obligatoria"}]


def test_import_invalid_start_date_format(actor_client):
    upload = _upload([HEADER, _row(start="no-es-fecha")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [
        {"row": 2, "error": "formato de fecha inválido, usar YYYY-MM-DD"}
    ]


def test_import_missing_end_date(actor_client):
    upload = _upload([HEADER, _row(end="")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [{"row": 2, "error": "la fecha de fin es obligatoria"}]


def test_import_invalid_hours(actor_client):
    upload = _upload([HEADER, _row(hours="no-numero")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [{"row": 2, "error": "Tiempo objetivo inválido"}]


def test_import_unknown_email(actor_client):
    upload = _upload([HEADER, _row(email="nadie@example.com")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert response.data["errors"] == [
        {"row": 2, "error": 'Usuario no encontrado: "nadie@example.com"'}
    ]


# --- import: normalización de `type` ---------------------------------------


def test_import_normalizes_type_case_insensitive(actor_client):
    upload = _upload([HEADER, _row(type="seguimiento")])
    actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert Task.objects.get(title="Informe").type == "SEGUIMIENTO"


def test_import_invalid_type_defaults_to_fija(actor_client):
    upload = _upload([HEADER, _row(type="NO_ES_UN_TIPO")])
    actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    assert Task.objects.get(title="Informe").type == "FIJA"


# --- import: fechas ambiguas y celdas de fecha nativas ----------------------


def test_import_accepts_dd_mm_yyyy_slash_format(actor_client):
    upload = _upload([HEADER, _row(start="01/07/2026", end="15/07/2026")])
    actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")
    task = Task.objects.get(title="Informe")
    assert task.start_date.date() == date(2026, 7, 1)
    assert task.end_date.date() == date(2026, 7, 15)


def test_import_accepts_native_date_cells(actor_client):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(HEADER)
    sheet.append(
        ["Informe", "d", "ALTA", "MENSUAL", date(2026, 7, 1), date(2026, 7, 15), 8, "", "FIJA"]
    )
    buffer = BytesIO()
    workbook.save(buffer)
    upload = SimpleUploadedFile("tasks.xlsx", buffer.getvalue(), content_type=XLSX_CONTENT_TYPE)

    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")

    assert response.data == {"imported": 1, "errors": []}
    task = Task.objects.get(title="Informe")
    assert task.start_date.date() == date(2026, 7, 1)
    assert task.end_date.date() == date(2026, 7, 15)


# --- import: una fila fallida no bloquea el resto del batch -----------------


def test_import_continues_after_row_creation_failure(actor_client, monkeypatch):
    from apps.tasks import services as tasks_services

    original_create = tasks_services.Task.objects.create

    def flaky_create(**kwargs):
        if kwargs["title"] == "Falla":
            raise ValueError("boom")
        return original_create(**kwargs)

    monkeypatch.setattr(tasks_services.Task.objects, "create", flaky_create)

    upload = _upload([HEADER, _row(title="Falla"), _row(title="Exito")])
    response = actor_client.post("/api/v1/tasks/import/", {"file": upload}, format="multipart")

    assert response.data == {
        "imported": 1,
        "errors": [{"row": 2, "error": "Error al crear la tarea"}],
    }
    assert Task.objects.filter(title="Exito").exists()
    assert not Task.objects.filter(title="Falla").exists()
