"""Captura masiva de Horas extras (guardado por lote), exportación completa y
restricción de unicidad en base de datos."""

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from django.db import IntegrityError, transaction

from apps.novedades.excel import OVERTIME_EXPORT_HEADERS, OVERTIME_HEADERS
from apps.novedades.models import NovedadEvent, OvertimeRecord
from apps.novedades.periods import PeriodKey
from apps.novedades.services import (
    CollaboratorInput,
    CollaboratorService,
    NovedadesError,
    OvertimeService,
    PeriodService,
)

from .conftest import access

pytestmark = pytest.mark.django_db

OCT = PeriodKey(2026, 10)
MON = date(2026, 10, 12)
TUE = date(2026, 10, 13)
SAT = date(2026, 10, 10)
SUN = date(2026, 10, 11)
BASE = "/api/v1/novedades"


def _collab(sup, cedula, name):
    return CollaboratorService.create(
        access(sup),
        CollaboratorInput(
            cedula=cedula, full_name=name, regional="Sierra", area="Bodega", cargo="Operador"
        ),
    )


def _row(key, collaborator, day=MON, **values):
    return {"client_key": key, "collaborator_id": collaborator.pk, "date": day, **values}


@pytest.fixture
def team(supervisor):
    return [_collab(supervisor, f"010203040{i}", f"Persona {i}") for i in range(1, 4)]


def test_saves_many_rows_in_one_operation_with_identity_from_database(supervisor, team):
    rows = [
        _row(f"r{i}", c, overtime_hours="2", work_done="Inventario") for i, c in enumerate(team)
    ]
    # Datos de identidad enviados desde el navegador se ignoran.
    rows[0]["cedula"] = "9999999999"
    result = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    assert result["summary"]["created"] == 3 and result["summary"]["errors"] == 0
    assert {r["status"] for r in result["rows"]} == {"created"}
    assert OvertimeRecord.objects.count() == 3
    assert set(OvertimeRecord.objects.values_list("collaborator__cedula", flat=True)) == {
        c.cedula for c in team
    }
    assert NovedadEvent.objects.filter(action="overtime.created").count() == 3


def test_invalid_rows_are_reported_per_field_and_valid_ones_are_saved(supervisor, team):
    rows = [
        _row("ok", team[0], overtime_hours="2"),
        _row("neg", team[1], overtime_hours="-1"),
        _row("txt", team[2], overtime_hours="abc", saturday_hours="1.234"),
    ]
    result = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    by_key = {r["client_key"]: r for r in result["rows"]}
    assert by_key["ok"]["status"] == "created"
    assert by_key["neg"]["status"] == "error" and "overtime_hours" in by_key["neg"]["errors"]
    assert set(by_key["txt"]["errors"]) == {"overtime_hours", "saturday_hours"}
    assert OvertimeRecord.objects.count() == 1
    assert result["summary"] == {
        "created": 1,
        "updated": 0,
        "unchanged": 0,
        "errors": 2,
        "total": 3,
    }


def test_retry_of_same_batch_does_not_duplicate(supervisor, team):
    rows = [_row(f"r{i}", c, overtime_hours="2") for i, c in enumerate(team)]
    first = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    retry = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    assert OvertimeRecord.objects.count() == 3
    assert {r["status"] for r in retry["rows"]} == {"unchanged"}
    assert [r["record_id"] for r in retry["rows"]] == [r["record_id"] for r in first["rows"]]


def test_new_row_against_different_existing_record_is_a_duplicate(supervisor, team):
    OvertimeService.bulk_save(access(supervisor), OCT, [_row("a", team[0], overtime_hours="2")])
    result = OvertimeService.bulk_save(
        access(supervisor), OCT, [_row("b", team[0], overtime_hours="5")]
    )
    assert result["rows"][0]["status"] == "duplicate"
    assert OvertimeRecord.objects.get().overtime_hours == Decimal("2.00")


def test_duplicates_inside_the_same_batch(supervisor, team):
    rows = [_row("a", team[0], overtime_hours="2"), _row("b", team[0], overtime_hours="3")]
    result = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    assert [r["status"] for r in result["rows"]] == ["created", "duplicate"]


def test_editing_existing_rows_updates_without_copies_and_allows_date_change(supervisor, team):
    saved = OvertimeService.bulk_save(
        access(supervisor), OCT, [_row("a", team[0], overtime_hours="2")]
    )
    rid, version = saved["rows"][0]["record_id"], saved["rows"][0]["version"]
    edit = {"client_key": "a", "id": rid, "version": version, "date": TUE, "overtime_hours": "4"}
    result = OvertimeService.bulk_save(access(supervisor), OCT, [edit])
    assert result["rows"][0]["status"] == "updated"
    record = OvertimeRecord.objects.get()
    assert (record.date, record.overtime_hours, record.version) == (TUE, Decimal("4.00"), 2)
    # Versión vieja: conflicto, no se pisa.
    stale = OvertimeService.bulk_save(access(supervisor), OCT, [{**edit, "overtime_hours": "6"}])
    assert stale["rows"][0]["status"] == "conflict"


def test_weekends_and_holidays_are_allowed_in_overtime(supervisor, team):
    from apps.configuration.models import Holiday

    Holiday.objects.create(date=TUE, name="Feriado de prueba", year=2026)
    rows = [
        _row("sat", team[0], SAT, saturday_hours="4"),
        _row("sun", team[1], SUN, overtime_hours="3"),
        _row("hol", team[2], TUE, overtime_hours="2"),
    ]
    result = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    assert {r["status"] for r in result["rows"]} == {"created"}
    flags = {
        r.date: (r.is_saturday, r.is_sunday, r.is_holiday) for r in OvertimeRecord.objects.all()
    }
    assert flags == {
        SAT: (True, False, False),
        SUN: (False, True, False),
        TUE: (False, False, True),
    }


def test_out_of_scope_and_closed_period_rows_are_locked(supervisor, supervisor2, team):
    other = _collab(supervisor2, "0911111111", "Ajena")
    result = OvertimeService.bulk_save(
        access(supervisor),
        OCT,
        [_row("mine", team[0], overtime_hours="1"), _row("other", other, overtime_hours="1")],
    )
    assert [r["status"] for r in result["rows"]] == ["created", "locked"]
    PeriodService.close(access(supervisor), OCT, supervisor)
    closed = OvertimeService.bulk_save(
        access(supervisor), OCT, [_row("x", team[1], overtime_hours="1")]
    )
    assert closed["rows"][0]["status"] == "locked"
    assert OvertimeRecord.objects.count() == 1


def test_dates_outside_period_and_missing_fields(supervisor, team):
    rows = [
        _row("out", team[0], date(2026, 10, 25), overtime_hours="1"),
        {
            "client_key": "nodate",
            "collaborator_id": team[1].pk,
            "date": None,
            "overtime_hours": "1",
        },
        _row("nohours", team[2]),
    ]
    result = OvertimeService.bulk_save(access(supervisor), OCT, rows)
    assert [r["status"] for r in result["rows"]] == ["error", "error", "error"]


def test_assistant_and_batch_limit(assistant, supervisor, team):
    with pytest.raises(NovedadesError):
        OvertimeService.bulk_save(access(assistant), OCT, [_row("a", team[0], overtime_hours="1")])
    big = [_row(str(i), team[0], overtime_hours="1") for i in range(OvertimeService.BULK_LIMIT + 1)]
    with pytest.raises(NovedadesError) as exc:
        OvertimeService.bulk_save(access(supervisor), OCT, big)
    assert exc.value.code == "batch_too_large"


def test_database_rejects_two_active_records_for_same_person_and_date(supervisor, team):
    OvertimeService.bulk_save(access(supervisor), OCT, [_row("a", team[0], overtime_hours="2")])
    record = OvertimeRecord.objects.get()
    with pytest.raises(IntegrityError), transaction.atomic():
        OvertimeRecord.objects.create(
            collaborator=team[0],
            period=record.period,
            supervisor=supervisor,
            date=MON,
            overtime_hours=1,
            created_by=supervisor,
            updated_by=supervisor,
        )


# ── API ───────────────────────────────────────────────────────────────────────


def test_bulk_endpoint_returns_per_row_results(client_for, supervisor, team):
    payload = {
        "period": "2026-10",
        "rows": [
            {
                "client_key": "a",
                "collaborator_id": team[0].pk,
                "date": "2026-10-12",
                "overtime_hours": "2",
            },
            {
                "client_key": "b",
                "collaborator_id": team[1].pk,
                "date": "2026-10-12",
                "overtime_hours": "x",
            },
        ],
    }
    response = client_for(supervisor).post(f"{BASE}/overtime/bulk/", payload, format="json")
    assert response.status_code == 200
    body = response.json()
    assert [r["status"] for r in body["rows"]] == ["created", "error"]


def test_consolidated_pagination_keeps_totals_and_export_is_complete(
    client_for, admin, supervisor, assistant, team
):
    rows = [_row(f"{c.pk}-{d}", c, d, overtime_hours="1") for c in team for d in (MON, TUE)]
    OvertimeService.bulk_save(access(supervisor), OCT, rows)
    page = client_for(assistant).get(f"{BASE}/overtime/?period=2026-10&limit=2&offset=0").json()
    assert len(page["records"]) == 2 and page["count"] == 6 and page["indicators"]["records"] == 6
    assert page["indicators"]["overtime_hours"] == "6"
    assert page["indicators"]["supervisors_pending"] == 1
    export = client_for(assistant).get(f"{BASE}/overtime/export/?period=2026-10")
    ws = openpyxl.load_workbook(io.BytesIO(export.content)).active
    assert ws.max_row == 7  # encabezado + 6, aunque la pantalla pagine
    assert [c.value for c in ws[1]] == OVERTIME_EXPORT_HEADERS
    assert NovedadEvent.objects.filter(action="export.overtime").exists()


def test_export_keeps_leading_zeros_and_neutralizes_formulas(client_for, supervisor, assistant):
    collaborator = _collab(supervisor, "0012345678", "=CMD|' /C calc'!A0")
    OvertimeService.bulk_save(
        access(supervisor),
        OCT,
        [_row("a", collaborator, overtime_hours="1", observations="+SUM(A1)")],
    )
    content = client_for(assistant).get(f"{BASE}/overtime/export/?period=2026-10").content
    ws = openpyxl.load_workbook(io.BytesIO(content)).active
    assert ws.cell(row=2, column=2).value == "0012345678"
    assert ws.cell(row=2, column=2).number_format == "@"
    assert ws.cell(row=2, column=3).value.startswith("'=")
    assert ws.cell(row=2, column=14).value == "'+SUM(A1)"


def test_overtime_template_has_the_13_headers(client_for, supervisor):
    response = client_for(supervisor).get(f"{BASE}/overtime/template/")
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in ws[1]] == OVERTIME_HEADERS and len(OVERTIME_HEADERS) == 13


def test_jornada_filter(client_for, supervisor, team):
    OvertimeService.bulk_save(
        access(supervisor),
        OCT,
        [
            _row("sat", team[0], SAT, saturday_hours="2"),
            _row("mon", team[1], MON, overtime_hours="1"),
        ],
    )
    body = client_for(supervisor).get(f"{BASE}/overtime/?period=2026-10&jornada=sabado").json()
    assert [r["cedula"] for r in body["records"]] == [team[0].cedula]
