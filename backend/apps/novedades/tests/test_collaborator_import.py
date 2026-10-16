"""Importación de la Base de colaboradores: carga inicial sin registros
previos, creación vs. actualización, inválidos, duplicados e idempotencia."""

import io
from datetime import date

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.notifications.models import Notification
from apps.novedades import access as perms
from apps.novedades.excel import COLLABORATOR_TEMPLATE_HEADERS
from apps.novedades.models import Collaborator, ImportBatch
from apps.novedades.services import CollaboratorInput, CollaboratorService, PermissionAdminService

from .conftest import access

pytestmark = pytest.mark.django_db
BASE = "/api/v1/novedades"
HEADER = ["CEDULA", "NOMBRES COMPLETOS", "REGIONAL", "ÁREA", "CARGO"]


def _xlsx(rows) -> SimpleUploadedFile:
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("base.xlsx", buffer.getvalue())


def _grant_import(admin, user):
    PermissionAdminService.create_override(
        access(admin),
        user_id=user.pk,
        codename=perms.IMPORTAR,
        effect="GRANT",
        reason="Carga inicial",
    )


def _preview(client, rows, **extra):
    return client.post(
        f"{BASE}/imports/preview/",
        {"kind": "COLLABORATORS", "file": _xlsx(rows), **extra},
        format="multipart",
    )


def _confirm(client, batch_id, strategy="skip"):
    return client.post(f"{BASE}/imports/{batch_id}/confirm/", {"strategy": strategy}, format="json")


def test_initial_load_creates_collaborators_that_did_not_exist(client_for, admin, supervisor):
    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    rows = [
        HEADER,
        ["0102030405", "Ana Pérez", "Sierra", "Bodega", "Operador"],
        [102030406, "Beto Ruiz", "Costa", "Rutas", "Mensajero"],  # número: perdió el 0 inicial
        ["1712345678", "  Carla   Mora ", "", "", ""],
    ]
    response = _preview(client, rows)
    assert response.status_code == 201, response.json()
    body = response.json()
    assert body["errors"] == []
    assert body["summary"]["new_items"] == 3 and body["summary"]["invalid_rows"] == 0
    assert not Collaborator.objects.exists()  # la vista previa no guarda nada

    confirmed = _confirm(client, body["id"]).json()
    assert confirmed["summary"]["created"] == 3
    cedulas = set(Collaborator.objects.values_list("cedula", flat=True))
    assert cedulas == {"0102030405", "0102030406", "1712345678"}
    assert Collaborator.objects.get(cedula="1712345678").full_name == "Carla Mora"
    assert all(c.current_supervisor == supervisor for c in Collaborator.objects.all())
    # Una sola notificación resumen al Administrador, no una por persona.
    # Alerta azul: resumen solo de lo creado, con origen importación.
    alerts = list(Notification.objects.filter(user=admin).values_list("message", "category"))
    assert len(alerts) == 1
    message, category = alerts[0]
    assert category == "new_collaborator"
    assert "registrado(s) en la base: 3" in message and "origen: importación" in message


def test_invalid_rows_and_duplicates_inside_the_file(client_for, admin, supervisor):
    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    rows = [
        HEADER,
        ["0102030405", "Ana", "", "", ""],
        ["0102030405", "Ana repetida", "", "", ""],
        ["", "Sin cédula", "", "", ""],
        ["12-AB#", "Caracteres raros", "", "", ""],
        ["0911111111", "", "", "", ""],
        ["0922222222", "Válida", "", "", ""],
    ]
    body = _preview(client, rows).json()
    assert body["summary"]["valid_rows"] == 2
    assert body["summary"]["duplicates_in_file"] == 1
    assert body["summary"]["invalid_rows"] == 4
    assert {e["row"] for e in body["errors"]} == {3, 4, 5, 6}
    _confirm(client, body["id"])
    assert Collaborator.objects.count() == 2


def test_existing_collaborators_skip_or_update_only_allowed_fields(client_for, admin, supervisor):
    _grant_import(admin, supervisor)
    collaborator = CollaboratorService.create(
        access(supervisor),
        CollaboratorInput(
            cedula="0102030405", full_name="Ana", regional="Sierra", cargo="Operador"
        ),
    )
    CollaboratorService.change_status(access(supervisor), collaborator.pk, "SUSPENDED", "Sanción")
    client = client_for(supervisor)
    rows = [
        HEADER,
        ["0102030405", "Ana Pérez", "Sierra", "Bodega", "Jefe"],
        ["0999999999", "Nuevo", "", "", ""],
    ]

    body = _preview(client, rows).json()
    assert body["summary"]["new_items"] == 1 and body["summary"]["to_update"] == 1
    existing = next(r for r in body["rows"] if r["status"] == "EXISTING")
    assert set(existing["changes"]) == {"full_name", "area", "cargo"}

    _confirm(client, body["id"], "skip")
    collaborator.refresh_from_db()
    assert collaborator.full_name == "Ana" and collaborator.cargo == "Operador"

    blue = Notification.objects.filter(category="new_collaborator")
    blue_before = blue.count()
    body = _preview(client, rows).json()
    assert body["summary"]["new_items"] == 0 and body["summary"]["to_update"] == 1
    result = _confirm(client, body["id"], "update").json()
    assert result["summary"]["updated"] == 1
    # Una actualización no es un ingreso nuevo: no genera alerta azul.
    assert blue.count() == blue_before
    collaborator.refresh_from_db()
    assert (collaborator.full_name, collaborator.area, collaborator.cargo) == (
        "Ana Pérez",
        "Bodega",
        "Jefe",
    )
    # Lo que una importación nunca toca: estado y supervisor.
    assert collaborator.status == "SUSPENDED" and collaborator.current_supervisor == supervisor


def test_collaborator_of_another_supervisor_is_rejected_not_reassigned(
    client_for, admin, supervisor, supervisor2
):
    _grant_import(admin, supervisor2)
    CollaboratorService.create(
        access(supervisor), CollaboratorInput(cedula="0102030405", full_name="Ana")
    )
    body = _preview(client_for(supervisor2), [HEADER, ["0102030405", "Ana", "", "", ""]]).json()
    assert body["summary"]["valid_rows"] == 0
    assert "ya pertenece a la base" in body["errors"][0]["message"]
    _confirm(client_for(supervisor2), body["id"], "update")
    assert Collaborator.objects.get().current_supervisor == supervisor


def test_reimport_and_reconfirm_do_not_duplicate(client_for, admin, supervisor):
    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    rows = [HEADER, ["0102030405", "Ana", "", "", ""]]
    first = _preview(client, rows).json()
    _confirm(client, first["id"])
    assert _confirm(client, first["id"]).status_code == 200  # confirmar dos veces
    second = _preview(client, rows).json()
    assert second["summary"]["new_items"] == 0 and second["summary"]["unchanged"] == 1
    _confirm(client, second["id"], "update")
    assert Collaborator.objects.count() == 1


def test_collaborator_created_between_preview_and_confirm_is_not_duplicated(
    client_for, admin, supervisor
):
    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    body = _preview(client, [HEADER, ["0102030405", "Ana", "", "", ""]]).json()
    CollaboratorService.create(
        access(supervisor), CollaboratorInput(cedula="0102030405", full_name="Ana")
    )
    result = _confirm(client, body["id"]).json()
    assert result["summary"]["created"] == 0
    assert Collaborator.objects.count() == 1


def test_accepts_the_worked_days_file_format_as_base(client_for, admin, supervisor):
    """La planilla de Días laborados que la operación ya usa (columna
    "Empleado", fechas) sirve para cargar la base: las fechas se ignoran."""
    _grant_import(admin, supervisor)
    rows = [
        ["Supervisor", "Cargo", "Cédula", "Empleado", date(2026, 10, 12)],
        ["Sup", "Operador", "0102030405", "Ana Pérez", 8],
    ]
    body = _preview(client_for(supervisor), rows).json()
    assert body["summary"]["new_items"] == 1
    assert "Supervisor" in body["summary"]["ignored_columns"]


def test_scope_and_permissions(client_for, admin, supervisor, supervisor2, assistant):
    # Sin facultad de importar: 403.
    assert (
        _preview(client_for(supervisor), [HEADER, ["0102030405", "Ana", "", "", ""]]).status_code
        == 403
    )
    assert (
        _preview(client_for(assistant), [HEADER, ["0102030405", "Ana", "", "", ""]]).status_code
        == 403
    )
    # Un supervisor siempre importa a SU base, aunque pida otra.
    _grant_import(admin, supervisor)
    body = _preview(
        client_for(supervisor),
        [HEADER, ["0102030405", "Ana", "", "", ""]],
        supervisor_id=supervisor2.pk,
    ).json()
    assert body["summary"]["supervisor_id"] == supervisor.pk
    # El Administrador debe elegir el supervisor destino.
    assert (
        _preview(client_for(admin), [HEADER, ["0102030405", "Ana", "", "", ""]]).status_code == 400
    )
    ok = _preview(
        client_for(admin),
        [HEADER, ["0911111111", "Beto", "", "", ""]],
        supervisor_id=supervisor2.pk,
    )
    assert ok.status_code == 201 and ok.json()["summary"]["supervisor_id"] == supervisor2.pk


def test_missing_required_columns_and_template(client_for, admin, supervisor):
    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    response = _preview(client, [["REGIONAL", "CARGO"], ["Sierra", "Operador"]])
    assert response.status_code == 400 and "CEDULA" in response.json()["error"]["message"]
    template = client.get(f"{BASE}/collaborators/template/")
    ws = openpyxl.load_workbook(io.BytesIO(template.content)).active
    assert [c.value for c in ws[1]] == COLLABORATOR_TEMPLATE_HEADERS
    assert ws.cell(row=2, column=1).number_format == "@"
    assert ImportBatch.objects.count() == 0


# ── Columna NUEVO INGRESO (2026-10-16) ────────────────────────────────────────
# X (sin distinguir mayúsculas ni espacios) = nuevo ingreso; vacía = no; otro
# valor = fila rechazada. Solo clasifica ALTAS: en un existente no cambia nada.
# El correo nunca sale al importar: el ingreso queda en espera del 2.º día.

NEW_HIRE_HEADER = [*HEADER, "NUEVO INGRESO"]


def test_template_has_the_exact_new_hire_header():
    from apps.novedades.excel import CollaboratorImportService

    ws = openpyxl.load_workbook(io.BytesIO(CollaboratorImportService.template())).active
    assert [c.value for c in ws[1]][-1] == "NUEVO INGRESO"


def test_new_hire_column_values(client_for, admin, supervisor):
    from django.core import mail

    from apps.novedades.models import CollaboratorEntryNotification

    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    rows = [
        NEW_HIRE_HEADER,
        ["0102030401", "Ana X Mayus", "", "", "", "X"],
        ["0102030402", "Beto x minus", "", "", "", "x"],
        ["0102030403", "Carla Espacios", "", "", "", "  x  "],
        ["0102030404", "Dani Vacia", "", "", "", None],
        ["0102030405", "Eva Solo Espacios", "", "", "", "   "],
        ["0102030406", "Fede Si", "", "", "", "Sí"],
        ["0102030407", "Gina Uno", "", "", "", 1],
        ["0102030408", "Hugo XX", "", "", "", "XX"],
    ]
    preview = _preview(client, rows)
    assert preview.status_code == 201, preview.json()
    body = preview.json()
    bad = {e["row"]: e for e in body["errors"]}
    # Inválidos: rechazados con mensaje claro, nunca tomados como nuevo ingreso.
    assert set(bad) == {7, 8, 9}
    assert all(e["column"] == "NUEVO INGRESO" and "escriba X" in e["message"] for e in bad.values())
    assert body["summary"]["new_items"] == 5 and body["summary"]["new_hires"] == 3

    confirmed = _confirm(client, body["id"]).json()["summary"]
    assert confirmed["created"] == 5 and confirmed["new_hires"] == 3
    marked = set(
        CollaboratorEntryNotification.objects.values_list("collaborator__cedula", flat=True)
    )
    assert marked == {"0102030401", "0102030402", "0102030403"}
    assert set(CollaboratorEntryNotification.objects.values_list("status", flat=True)) == {
        "WAITING"
    }
    assert not Collaborator.objects.filter(
        cedula__in=["0102030406", "0102030407", "0102030408"]
    ).exists()
    assert not [m for m in mail.outbox if "ingreso" in m.subject]  # nunca al importar


def test_reimport_never_changes_existing_classification(client_for, admin, supervisor):
    from apps.novedades.models import CollaboratorEntryNotification

    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    first = _preview(
        client,
        [
            NEW_HIRE_HEADER,
            ["0102030401", "Ana Nueva", "", "", "", "X"],
            ["0102030402", "Beto Antiguo", "", "", "", None],
        ],
    ).json()
    _confirm(client, first["id"])
    ana_event = CollaboratorEntryNotification.objects.get(collaborator__cedula="0102030401")

    # Reimportar: Ana con celda vacía NO pierde su marca; Beto con X NO la gana.
    again = _preview(
        client,
        [
            NEW_HIRE_HEADER,
            ["0102030401", "Ana Nueva", "Sierra", "", "", None],
            ["0102030402", "Beto Antiguo", "", "", "", "X"],
        ],
    ).json()
    assert again["summary"]["new_hires"] == 0
    warnings = again["summary"]["warnings"]
    assert [w["row"] for w in warnings] == [3] and "no cambia su clasificación" in warnings[0][
        "message"
    ]
    for strategy in ("update", "update"):  # confirmar dos veces tampoco duplica
        _confirm(client, again["id"], strategy)
    assert list(CollaboratorEntryNotification.objects.values_list("pk", flat=True)) == [
        ana_event.pk
    ]
    assert (
        Collaborator.objects.get(cedula="0102030401").regional == "Sierra"
    )  # el resto sí se actualiza


def test_files_without_the_column_still_import_as_not_new_hire(client_for, admin, supervisor):
    from apps.novedades.models import CollaboratorEntryNotification

    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    body = _preview(client, [HEADER, ["0102030401", "Ana Sin Columna", "", "", ""]]).json()
    assert body["errors"] == [] and body["summary"]["new_hires"] == 0
    _confirm(client, body["id"])
    assert Collaborator.objects.filter(cedula="0102030401").exists()
    assert not CollaboratorEntryNotification.objects.exists()


def test_imported_new_hire_mails_on_second_valid_day(client_for, admin, supervisor):
    from django.core import mail

    from apps.novedades.models import (
        CollaboratorEntryNotification,
        CollaboratorNotificationRecipient,
    )
    from apps.novedades.periods import PeriodKey
    from apps.novedades.services import WorkedDayService

    CollaboratorNotificationRecipient.objects.create(
        kind="ENTRY", email="ops@example.com", created_by=admin, updated_by=admin
    )
    _grant_import(admin, supervisor)
    client = client_for(supervisor)
    body = _preview(
        client, [NEW_HIRE_HEADER, ["0102030401", "Ana Importada", "", "", "", "X"]]
    ).json()
    _confirm(client, body["id"])
    ana = Collaborator.objects.get(cedula="0102030401")
    for day in (date(2026, 10, 8), date(2026, 10, 9)):
        WorkedDayService.save_cells(
            access(supervisor),
            PeriodKey(2026, 10),
            supervisor,
            [{"collaborator_id": ana.pk, "date": day, "hours": "8"}],
        )
    assert CollaboratorEntryNotification.objects.get(collaborator=ana).status == "SENT"
    assert [m.to for m in mail.outbox if "ingreso" in m.subject] == [["ops@example.com"]]
