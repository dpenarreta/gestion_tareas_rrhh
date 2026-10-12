"""Superficie HTTP: control de acceso en el servidor, manipulación de
parámetros, importación/exportación de Excel."""

import io
from datetime import date

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.novedades import access as perms
from apps.novedades.excel import OVERTIME_EXPORT_HEADERS, OVERTIME_HEADERS
from apps.novedades.models import Collaborator, ImportBatch, OvertimeRecord, WorkedDay
from apps.novedades.services import PermissionAdminService

from .conftest import access

pytestmark = pytest.mark.django_db
BASE = "/api/v1/novedades"


def _create(client, cedula="0102030405", name="Ana Pérez", **extra):
    payload = {
        "cedula": cedula,
        "full_name": name,
        "regional": "Sierra",
        "area": "Bodega",
        "cargo": "Operador",
    }
    payload.update(extra)
    return client.post(f"{BASE}/collaborators/", payload, format="json")


def _xlsx(rows) -> SimpleUploadedFile:
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return SimpleUploadedFile("datos.xlsx", buffer.getvalue())


def test_users_without_module_get_403_and_it_is_audited(client_for, other_user):
    from apps.core.models import AuditLog

    response = client_for(other_user).get(f"{BASE}/me/")
    assert response.status_code == 403
    assert AuditLog.objects.filter(
        actor=other_user, action="access_denied", module="novedades"
    ).exists()


def test_supervisor_cannot_read_other_team_via_params(client_for, supervisor, supervisor2):
    assert _create(client_for(supervisor)).status_code == 201
    other = client_for(supervisor2)
    assert other.get(f"{BASE}/worked-days/?supervisor={supervisor.pk}").status_code == 403
    assert other.get(f"{BASE}/collaborators/").json() == []
    assert other.get(f"{BASE}/overtime/?supervisor={supervisor.pk}").json()["records"] == []


def test_supervisor_cannot_export_without_permission_even_calling_api(
    client_for, admin, supervisor
):
    client = client_for(supervisor)
    assert client.get(f"{BASE}/overtime/export/").status_code == 403
    PermissionAdminService.set_role_permission(
        access(admin),
        role=perms.SUPERVISOR_ROLE,
        codename=perms.EXPORTAR,
        enabled=True,
        reason="ok",
    )
    assert client_for(supervisor).get(f"{BASE}/overtime/export/").status_code == 200


def test_revoked_access_takes_effect_without_relogin(client_for, admin, supervisor):
    client = client_for(supervisor)
    assert client.get(f"{BASE}/me/").status_code == 200
    PermissionAdminService.create_override(
        access(admin), user_id=supervisor.pk, codename=perms.ACCEDER, effect="DENY", reason="Baja"
    )
    assert client.get(f"{BASE}/me/").status_code == 403


def test_worked_days_endpoint_rejects_weekend_cells(client_for, supervisor):
    client = client_for(supervisor)
    cid = _create(client).json()["id"]
    response = client.post(
        f"{BASE}/worked-days/save/",
        {
            "period": "2026-10",
            "cells": [{"collaborator_id": cid, "date": "2026-10-10", "hours": "8"}],
        },
        format="json",
    )
    assert response.status_code == 400
    assert "sábado y domingo" in response.json()["error"]["details"]["cells"][0]["message"]


def test_overtime_export_has_14_headers_with_supervisor_seventh(
    client_for, admin, supervisor, assistant
):
    sup_client = client_for(supervisor)
    cid = _create(sup_client).json()["id"]
    sup_client.post(
        f"{BASE}/overtime/",
        {
            "collaborator_id": cid,
            "date": "2026-10-12",
            "overtime_hours": "2",
            "work_done": "Inventario",
        },
        format="json",
    )
    response = client_for(assistant).get(f"{BASE}/overtime/export/?period=2026-10")
    assert response.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    headers = [c.value for c in ws[1]]
    assert headers == OVERTIME_EXPORT_HEADERS and len(headers) == 14 and headers[6] == "SUPERVISOR"
    data = [c.value for c in ws[2]]
    assert data[1] == "0102030405" and data[6] == "Sup1" and data[12] == "Inventario"
    # Filtros activos respetados.
    empty = client_for(assistant).get(f"{BASE}/overtime/export/?period=2026-10&cedula=999")
    assert openpyxl.load_workbook(io.BytesIO(empty.content)).active.max_row == 1


def test_overtime_list_has_13_columns_payload_and_assistant_cannot_write(
    client_for, supervisor, assistant
):
    cid = _create(client_for(supervisor)).json()["id"]
    response = client_for(assistant).post(
        f"{BASE}/overtime/",
        {"collaborator_id": cid, "date": "2026-10-12", "overtime_hours": "2"},
        format="json",
    )
    assert response.status_code == 403  # el Asistente nunca registra novedades
    assert len(OVERTIME_HEADERS) == 13


def test_import_worked_days_preview_confirm_idempotent(client_for, admin, supervisor):
    PermissionAdminService.create_override(
        access(admin),
        user_id=supervisor.pk,
        codename=perms.IMPORTAR,
        effect="GRANT",
        reason="Carga",
    )
    client = client_for(supervisor)
    _create(client)
    rows = [
        [
            "Supervisor",
            "Cargo",
            "Corte anterior",
            "Cédula",
            "Empleado",
            date(2026, 10, 12),
            date(2026, 10, 10),
        ],
        ["Sup1", "Operador", "", "0102030405", "Ana Pérez", 8, None],
        ["Sup1", "Operador", "", "0999999999", "Nadie", 8, None],
    ]
    preview = client.post(
        f"{BASE}/imports/preview/",
        {"kind": "WORKED_DAYS", "period": "2026-10", "file": _xlsx(rows)},
        format="multipart",
    )
    assert preview.status_code == 201, preview.json()
    body = preview.json()
    assert body["summary"]["valid_rows"] == 1 and body["summary"]["invalid_rows"] == 1
    assert not WorkedDay.objects.exists()  # la vista previa no escribe nada
    confirm = client.post(
        f"{BASE}/imports/{body['id']}/confirm/", {"strategy": "skip"}, format="json"
    )
    assert confirm.status_code == 200 and confirm.json()["summary"]["applied"] == 1
    again = client.post(
        f"{BASE}/imports/{body['id']}/confirm/", {"strategy": "skip"}, format="json"
    )
    assert again.status_code == 200
    assert WorkedDay.objects.count() == 1

    # Reimportar el mismo archivo no duplica y se advierte.
    rows2 = [rows[0], rows[1]]
    preview2 = client.post(
        f"{BASE}/imports/preview/",
        {"kind": "WORKED_DAYS", "period": "2026-10", "file": _xlsx(rows2)},
        format="multipart",
    ).json()
    assert preview2["summary"]["existing_items"] == 1
    client.post(f"{BASE}/imports/{preview2['id']}/confirm/", {"strategy": "skip"}, format="json")
    assert WorkedDay.objects.count() == 1
    errors = client.get(f"{BASE}/imports/{body['id']}/errors/")
    assert errors.status_code == 200


def test_import_rejects_weekend_values(client_for, admin, supervisor):
    PermissionAdminService.create_override(
        access(admin),
        user_id=supervisor.pk,
        codename=perms.IMPORTAR,
        effect="GRANT",
        reason="Carga",
    )
    client = client_for(supervisor)
    _create(client)
    rows = [
        ["Supervisor", "Cargo", "Corte anterior", "Cédula", "Empleado", date(2026, 10, 10)],
        ["Sup1", "Operador", "", "0102030405", "Ana Pérez", 8],
    ]
    body = client.post(
        f"{BASE}/imports/preview/",
        {"kind": "WORKED_DAYS", "period": "2026-10", "file": _xlsx(rows)},
        format="multipart",
    ).json()
    assert body["summary"]["invalid_rows"] == 1
    assert "sábado y domingo" in body["errors"][0]["message"]


def test_import_overtime_validates_headers_and_duplicates(client_for, admin, supervisor):
    client = client_for(admin)
    _create(client, supervisor_id=supervisor.pk)
    bad = client.post(
        f"{BASE}/imports/preview/",
        {"kind": "OVERTIME", "period": "2026-10", "file": _xlsx([["FECHA", "CEDULA"]])},
        format="multipart",
    )
    assert bad.status_code == 400
    row = [date(2026, 10, 12), "0102030405", "Ana", "", "", "", 2, 0, "NO", "NO", "NO", "Carga", ""]
    body = client.post(
        f"{BASE}/imports/preview/",
        {"kind": "OVERTIME", "period": "2026-10", "file": _xlsx([OVERTIME_HEADERS, row, row])},
        format="multipart",
    ).json()
    assert body["summary"]["valid_rows"] == 1 and body["summary"]["invalid_rows"] == 1
    client.post(f"{BASE}/imports/{body['id']}/confirm/", {"strategy": "skip"}, format="json")
    assert OvertimeRecord.objects.count() == 1
    assert ImportBatch.objects.get(pk=body["id"]).status == "CONFIRMED"


def test_admin_national_scope_and_collaborator_creation_for_supervisor(
    client_for, admin, supervisor
):
    client = client_for(admin)
    assert _create(client).status_code == 400  # debe indicar el supervisor
    assert _create(client, supervisor_id=supervisor.pk).status_code == 201
    assert Collaborator.objects.get().current_supervisor == supervisor
    status = client.get(f"{BASE}/periods/status/?period=2026-10").json()
    assert status[0]["supervisor_id"] == supervisor.pk
