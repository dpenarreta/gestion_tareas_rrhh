"""Plantilla e importación de Días laborados: encabezados con las fechas
reales del período, ida y vuelta plantilla → importación y correspondencia
cédula ↔ nombre."""

import io
from datetime import date

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.novedades import access as perms
from apps.novedades.models import WorkedDay
from apps.novedades.periods import PeriodKey
from apps.novedades.services import (
    CollaboratorInput,
    CollaboratorService,
    PeriodService,
    PermissionAdminService,
)

from .conftest import access

pytestmark = pytest.mark.django_db
BASE = "/api/v1/novedades"
LONG_NAME = "ACOSTA ACOSTA MARCO ANTONIO DE LA CRUZ VILLAVICENCIO"


@pytest.fixture
def setup(client_for, admin, supervisor):
    PermissionAdminService.create_override(
        access(admin),
        user_id=supervisor.pk,
        codename=perms.IMPORTAR,
        effect="GRANT",
        reason="Carga",
    )
    collaborator = CollaboratorService.create(
        access(supervisor),
        CollaboratorInput(cedula="0102030405", full_name=LONG_NAME, cargo="Operador"),
    )
    return client_for(supervisor), collaborator


def _upload(client, wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return client.post(
        f"{BASE}/imports/preview/",
        {
            "kind": "WORKED_DAYS",
            "period": "2026-10",
            "file": SimpleUploadedFile("dias.xlsx", buffer.getvalue()),
        },
        format="multipart",
    )


def _sheet(rows):
    wb = openpyxl.Workbook()
    for row in rows:
        wb.active.append(row)
    return wb


def test_matrix_returns_full_official_name(setup):
    client, _ = setup
    matrix = client.get(f"{BASE}/worked-days/?period=2026-10").json()
    assert matrix["rows"][0]["full_name"] == LONG_NAME


def test_template_headers_are_the_exact_period_dates_as_text(setup):
    client, _ = setup
    response = client.get(f"{BASE}/worked-days/template/?period=2026-10")
    wb = openpyxl.load_workbook(io.BytesIO(response.content))
    ws = wb["Días laborados"]
    headers = [c.value for c in ws[1]]
    expected_dates = [f"{d:%d/%m/%Y}" for d in PeriodKey(2026, 10).dates()]
    assert headers == ["Supervisor", "Cargo", "Cédula", "Empleado", *expected_dates]
    assert headers[4] == "23/09/2026" and headers[-1] == "22/10/2026"
    assert [c.value for c in ws[2]][:4] == ["Sup1", "Operador", "0102030405", LONG_NAME]
    assert all(v is None for v in [c.value for c in ws[2]][4:])  # plantilla vacía
    # Sábado 10/10/2026 resaltado en amarillo.
    sat_col = headers.index("10/10/2026") + 1
    assert ws.cell(row=1, column=sat_col).fill.fgColor.rgb.endswith("FFF59D")
    assert "Instrucciones" in wb.sheetnames
    text = " ".join(str(r[0].value) for r in wb["Instrucciones"].iter_rows())
    assert "HORAS" in text and "vacía" in text and "amarillo" in text


def test_filled_template_imports_back(setup):
    client, collaborator = setup
    content = client.get(f"{BASE}/worked-days/template/?period=2026-10").content
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb["Días laborados"]
    headers = [c.value for c in ws[1]]
    ws.cell(row=2, column=headers.index("12/10/2026") + 1).value = 8
    ws.cell(row=2, column=headers.index("13/10/2026") + 1).value = "7,5"
    body = _upload(client, wb).json()
    assert body["errors"] == [] and body["summary"]["valid_rows"] == 1
    client.post(f"{BASE}/imports/{body['id']}/confirm/", {"strategy": "skip"}, format="json")
    hours = dict(WorkedDay.objects.filter(collaborator=collaborator).values_list("date", "hours"))
    assert {d: float(h) for d, h in hours.items()} == {
        date(2026, 10, 12): 8.0,
        date(2026, 10, 13): 7.5,
    }


def test_rejects_dates_outside_period_and_repeated_dates(setup):
    client, _ = setup
    outside = _upload(
        client,
        _sheet(
            [
                ["Supervisor", "Cargo", "Cédula", "Empleado", "23/10/2026"],
                ["S", "C", "0102030405", LONG_NAME, 8],
            ]
        ),
    ).json()
    assert outside["summary"]["valid_rows"] == 0
    assert "no pertenece al período" in outside["errors"][0]["message"]
    repeated = _upload(
        client,
        _sheet(
            [
                ["Supervisor", "Cargo", "Cédula", "Empleado", "12/10/2026", date(2026, 10, 12)],
                ["S", "C", "0102030405", LONG_NAME, 8, 8],
            ]
        ),
    ).json()
    assert "repetida" in repeated["errors"][0]["message"]


def test_name_must_correspond_to_cedula(setup):
    client, _ = setup
    header = ["Supervisor", "Cargo", "Cédula", "Empleado", "12/10/2026"]
    wrong = _upload(client, _sheet([header, ["S", "C", "0102030405", "PEREZ JUAN", 8]])).json()
    assert wrong["summary"]["valid_rows"] == 0
    assert "Revise la cédula" in wrong["errors"][0]["message"]
    partial = _upload(client, _sheet([header, ["S", "C", "0102030405", "Acosta Marco", 8]])).json()
    assert partial["summary"]["valid_rows"] == 1
    assert any("no coincide exactamente" in w["message"] for w in partial["summary"]["warnings"])


def test_weekend_and_closed_period_rules_still_apply(setup, supervisor):
    client, _ = setup
    header = ["Supervisor", "Cargo", "Cédula", "Empleado", "10/10/2026"]
    weekend = _upload(client, _sheet([header, ["S", "C", "0102030405", LONG_NAME, 8]])).json()
    assert "sábado y domingo" in weekend["errors"][0]["message"]
    body = _upload(
        client,
        _sheet(
            [
                ["Supervisor", "Cargo", "Cédula", "Empleado", "12/10/2026"],
                ["S", "C", "0102030405", LONG_NAME, 8],
            ]
        ),
    ).json()
    PeriodService.close(access(supervisor), PeriodKey(2026, 10), supervisor)
    closed = client.post(
        f"{BASE}/imports/{body['id']}/confirm/", {"strategy": "skip"}, format="json"
    )
    assert closed.status_code == 409
    assert not WorkedDay.objects.exists()
