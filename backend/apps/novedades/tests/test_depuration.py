"""Horas depuradas y rol Jefe de Operaciones (pedido del 2026-10-13).

Pruebas críticas (bloqueantes):
1. Independencia real de las tres bases: modificar Horas depuradas no cambia
   la base de los supervisores ni el consolidado del Asistente (verificado
   contra la base de datos, no contra la pantalla).
2. Una sola copia por mes: crear/abrir de nuevo, solicitudes simultáneas y la
   restricción única de la base nunca producen una segunda copia."""

import io
from datetime import date
from decimal import Decimal

import openpyxl
import pytest
from django.db import IntegrityError, transaction

from apps.novedades.excel import OVERTIME_EXPORT_HEADERS
from apps.novedades.models import (
    DepuratedOvertime,
    DepurationChange,
    OvertimeDepuration,
    OvertimeRecord,
    WorkedDay,
)
from apps.novedades.periods import MonthKey, PeriodKey
from apps.novedades.services import (
    CollaboratorInput,
    CollaboratorService,
    DepurationService,
    NovedadesError,
    OvertimeService,
    PeriodService,
    WorkedDayService,
)

from .conftest import _user, access

pytestmark = pytest.mark.django_db
BASE = "/api/v1/novedades"
SEP = MonthKey(2026, 9)  # hoy es 12/10/2026: septiembre ya terminó
FIELDS = (
    "overtime_hours",
    "saturday_hours",
    "is_saturday",
    "is_sunday",
    "is_holiday",
    "work_done",
    "observations",
)


@pytest.fixture
def jefe(db):
    return _user("jefe", "JEFE_OPERACIONES")


def _collab(sup, cedula="0102030405", name="Ana Pérez"):
    return CollaboratorService.create(
        access(sup),
        CollaboratorInput(cedula=cedula, full_name=name, regional="Sierra", cargo="Operador"),
    )


def _sep_record(admin, c, day=date(2026, 9, 15), hours="10"):
    # Septiembre ya pasó su plazo por fecha: lo registra el Administrador.
    return OvertimeService.create(
        access(admin),
        {"collaborator_id": c.pk, "date": day, "overtime_hours": hours, "work_done": "Inventario"},
    )


def _snapshot(record_id):
    """Valores persistidos del registro original (respaldo 1 y 2)."""
    r = OvertimeRecord.objects.get(pk=record_id)
    return {
        f: getattr(r, f)
        for f in (
            *FIELDS,
            "version",
            "deleted_at",
            "updated_at",
            "date",
            "collaborator_id",
            "supervisor_id",
        )
    }


def _consolidated(assistant):
    """Consolidado del Asistente: su vista nacional de Horas extras del mes."""
    qs = OvertimeService.filter(OvertimeService.queryset_for(access(assistant)), {"period": SEP})
    return sorted((r.pk, r.overtime_hours, r.saturday_hours, r.work_done) for r in qs)


# ── PRUEBA CRÍTICA 1: independencia real de las tres bases ────────────────────


def test_critical_1_depuration_changes_only_the_depurated_base(admin, supervisor, assistant, jefe):
    c = _collab(supervisor)
    record = _sep_record(admin, c)
    original_before = _snapshot(record.pk)
    consolidated_before = _consolidated(assistant)
    assert original_before["overtime_hours"] == Decimal("10.00")
    assert consolidated_before == [(record.pk, Decimal("10.00"), Decimal("0.00"), "Inventario")]

    created = DepurationService.create(access(jefe), SEP)
    depuration_id = created["depuration"]["id"]
    row = DepuratedOvertime.objects.get(depuration_id=depuration_id, source_id=record.pk)
    assert row.orig_overtime_hours == Decimal("10.00") and row.overtime_hours == Decimal("10.00")

    # Dos ediciones: 10 → 9 → 8.
    for value, version in (("9", 1), ("8", 2)):
        result = DepurationService.save_rows(
            access(jefe),
            SEP,
            [{"id": row.pk, "version": version, "overtime_hours": value, "reason": "Ajuste"}],
        )
        assert result["rows"][0]["status"] == "updated"

    row.refresh_from_db()
    assert row.overtime_hours == Decimal("8.00")  # la base depurada cambió
    assert row.orig_overtime_hours == Decimal("10.00")  # la fotografía, no
    # Respaldo 1 (supervisores) y 2 (consolidado del Asistente): idénticos.
    assert _snapshot(record.pk) == original_before
    assert _consolidated(assistant) == consolidated_before
    # Historial: original 10, anterior 9, nuevo 8, responsable y motivo.
    last = DepurationChange.objects.filter(row=row).first()
    assert (last.field, last.original_value, last.old_value, last.new_value) == (
        "overtime_hours",
        "10",
        "9",
        "8",
    )
    assert last.actor == jefe and last.reason == "Ajuste"
    listed = {r["id"]: r for r in DepurationService.rows(access(jefe), SEP)}
    assert listed[row.pk]["difference_hours"] == "-2" and listed[row.pk]["modified"] is True


# ── PRUEBA CRÍTICA 2: una sola copia por mes ──────────────────────────────────


def test_critical_2_single_copy_per_month(client_for, admin, supervisor, jefe, monkeypatch):
    c = _collab(supervisor)
    _sep_record(admin, c)
    _sep_record(admin, _collab(supervisor, "0102030406", "Beto"), date(2026, 9, 16), "4")
    client = client_for(jefe)
    first = client.post(
        f"{BASE}/depuration/create/", {"period": "2026-09", "confirm": True}, format="json"
    ).json()
    depuration_id = first["depuration"]["id"]
    assert first["created"] is True and first["depuration"]["rows"] == 2
    row = DepuratedOvertime.objects.filter(depuration_id=depuration_id).first()
    DepurationService.save_rows(
        access(jefe), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "1"}]
    )

    # Volver a abrir / recargar / nueva sesión / repetir la acción de creación.
    for _ in range(2):
        again = (
            client_for(jefe)
            .post(
                f"{BASE}/depuration/create/", {"period": "2026-09", "confirm": True}, format="json"
            )
            .json()
        )
        assert again["created"] is False and again["depuration"]["id"] == depuration_id
        status = client_for(jefe).get(f"{BASE}/depuration/", {"period": "2026-09"}).json()
        assert status["depuration"]["id"] == depuration_id
    row.refresh_from_db()
    assert row.overtime_hours == Decimal("1.00")  # el ajuste se conservó

    # Dos solicitudes simultáneas: la segunda no ve la copia al verificar y
    # choca con la restricción única al escribir -> recibe la misma base.
    real_get = DepurationService.get
    calls = {"n": 0}

    def racing_get(month):
        calls["n"] += 1
        return None if calls["n"] == 1 else real_get(month)

    monkeypatch.setattr(DepurationService, "get", staticmethod(racing_get))
    raced = DepurationService.create(access(jefe), SEP)
    assert raced["created"] is False and raced["depuration"]["id"] == depuration_id

    # En la base: una sola copia del mes y una fila por registro de origen.
    assert OvertimeDepuration.objects.filter(year=2026, month=9).count() == 1
    assert DepuratedOvertime.objects.filter(depuration_id=depuration_id).count() == 2
    with pytest.raises(IntegrityError), transaction.atomic():
        OvertimeDepuration.objects.create(year=2026, month=9, created_by=jefe)


# ── B. Cambios posteriores en el origen ───────────────────────────────────────


def test_later_source_changes_do_not_overwrite_and_are_flagged(admin, supervisor, jefe):
    c = _collab(supervisor)
    record = _sep_record(admin, c)
    other = _sep_record(admin, _collab(supervisor, "0102030406", "Beto"), date(2026, 9, 16), "4")
    DepurationService.create(access(jefe), SEP)
    row = DepuratedOvertime.objects.get(source_id=record.pk)
    DepurationService.save_rows(
        access(jefe), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "7"}]
    )
    # El origen cambia y otro se elimina por las vías autorizadas.
    OvertimeService.update(access(admin), record.pk, {"overtime_hours": "12", "version": 1})
    OvertimeService.delete(access(admin), other.pk, "Duplicado")
    row.refresh_from_db()
    assert row.overtime_hours == Decimal("7.00") and row.orig_overtime_hours == Decimal("10.00")
    states = {r["source_id"]: r["source_state"] for r in DepurationService.rows(access(jefe), SEP)}
    assert states == {record.pk: "changed", other.pk: "deleted"}
    assert DepuratedOvertime.objects.filter(source_id=other.pk).exists()  # no se borró la copia


def test_excluding_a_row_never_deletes_the_source(admin, supervisor, jefe):
    record = _sep_record(admin, _collab(supervisor))
    DepurationService.create(access(jefe), SEP)
    row = DepuratedOvertime.objects.get()
    DepurationService.save_rows(access(jefe), SEP, [{"id": row.pk, "version": 1, "excluded": True}])
    record.refresh_from_db()
    assert record.deleted_at is None
    assert DepurationChange.objects.filter(
        field="excluded", original_value="No", new_value="Sí"
    ).exists()


# ── C. Permisos ───────────────────────────────────────────────────────────────


def test_jefe_reads_but_never_writes_original_bases(client_for, admin, supervisor, jefe):
    c = _collab(supervisor)
    record = _sep_record(admin, c)
    acc = access(jefe)
    assert (
        acc.national_read and not acc.can_write_for_supervisor(supervisor.pk) and acc.can_depurate
    )
    # Días laborados: consulta sí, edición no.
    matrix = WorkedDayService.matrix(acc, PeriodKey(2026, 10), supervisor)
    assert matrix["editable"] is False and matrix["closable"] is False
    with pytest.raises(NovedadesError):
        WorkedDayService.save_cells(
            acc,
            PeriodKey(2026, 10),
            supervisor,
            [{"collaborator_id": c.pk, "date": date(2026, 10, 9), "hours": "8"}],
        )
    # Horas extras originales: consulta sí, ninguna vía de escritura.
    assert OvertimeService.queryset_for(acc).filter(pk=record.pk).exists()
    for action in (
        lambda: OvertimeService.create(
            acc, {"collaborator_id": c.pk, "date": date(2026, 10, 9), "overtime_hours": "1"}
        ),
        lambda: OvertimeService.update(acc, record.pk, {"overtime_hours": "1", "version": 1}),
        lambda: OvertimeService.delete(acc, record.pk, "x"),
        lambda: OvertimeService.bulk_save(
            acc,
            MonthKey(2026, 10),
            [
                {
                    "client_key": "a",
                    "collaborator_id": c.pk,
                    "date": date(2026, 10, 9),
                    "overtime_hours": "1",
                }
            ],
        ),
        lambda: PeriodService.close(acc, PeriodKey(2026, 10), supervisor),
    ):
        with pytest.raises(NovedadesError):
            action()
    client = client_for(jefe)
    # Llamadas directas a la API: rechazadas.
    assert client.patch(
        f"{BASE}/overtime/{record.pk}/", {"overtime_hours": "1"}, format="json"
    ).status_code in (403, 404)
    assert client.delete(f"{BASE}/overtime/{record.pk}/", {}, format="json").status_code in (
        403,
        404,
    )
    assert client.get(f"{BASE}/history/").status_code == 403  # historial general: no
    assert client.get(f"{BASE}/permissions/config/").status_code == 403  # administración: no
    assert (
        client.patch(f"{BASE}/deadlines/settings/", {"worked_days": 72}, format="json").status_code
        == 403
    )
    assert OvertimeRecord.objects.get(pk=record.pk).overtime_hours == Decimal("10.00")
    assert not WorkedDay.objects.exists()


def test_depuration_access_by_role(client_for, admin, supervisor, assistant, jefe):
    _sep_record(admin, _collab(supervisor))
    DepurationService.create(access(jefe), SEP)
    row = DepuratedOvertime.objects.get()
    for user in (supervisor, assistant):
        assert client_for(user).get(f"{BASE}/depuration/", {"period": "2026-09"}).status_code == 403
    # Administrador: consulta y exporta, pero no edita.
    admin_client = client_for(admin)
    status = admin_client.get(f"{BASE}/depuration/", {"period": "2026-09"}).json()
    assert (
        status["exists"] is True and status["can_edit"] is False and status["can_create"] is False
    )
    assert admin_client.get(f"{BASE}/depuration/export/", {"period": "2026-09"}).status_code == 200
    with pytest.raises(NovedadesError):
        DepurationService.save_rows(
            access(admin), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "1"}]
        )
    # El mes en curso no se puede copiar todavía.
    response = client_for(jefe).post(
        f"{BASE}/depuration/create/", {"period": "2026-10", "confirm": True}, format="json"
    )
    assert response.status_code == 409
    # Sin base: estado informativo, sin errores.
    empty = admin_client.get(f"{BASE}/depuration/", {"period": "2026-08"}).json()
    assert empty["exists"] is False and "depuration" not in empty


def test_invalid_values_are_rejected_and_conflicts_detected(admin, supervisor, jefe):
    _sep_record(admin, _collab(supervisor))
    DepurationService.create(access(jefe), SEP)
    row = DepuratedOvertime.objects.get()
    bad = DepurationService.save_rows(
        access(jefe), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "-3"}]
    )
    assert bad["rows"][0]["status"] == "error"
    DepurationService.save_rows(
        access(jefe), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "5"}]
    )
    stale = DepurationService.save_rows(
        access(jefe), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "6"}]
    )
    assert stale["rows"][0]["status"] == "conflict"  # respuesta atrasada: no pisa
    retry = DepurationService.save_rows(
        access(jefe), SEP, [{"id": row.pk, "version": 1, "overtime_hours": "5"}]
    )
    assert retry["rows"][0]["status"] == "unchanged"  # reintento de lo ya guardado
    row.refresh_from_db()
    assert row.overtime_hours == Decimal("5.00") and row.version == 2


# ── F. Exportación ────────────────────────────────────────────────────────────


def test_export_contains_only_saved_depurated_values(client_for, admin, supervisor, jefe):
    c = _collab(supervisor)
    keep = _sep_record(admin, c, date(2026, 9, 12))  # sábado
    drop = _sep_record(admin, _collab(supervisor, "0102030406", "Beto"), date(2026, 9, 16), "4")
    DepurationService.create(access(jefe), SEP)
    kept = DepuratedOvertime.objects.get(source_id=keep.pk)
    dropped = DepuratedOvertime.objects.get(source_id=drop.pk)
    DepurationService.save_rows(
        access(jefe),
        SEP,
        [
            {"id": kept.pk, "version": 1, "overtime_hours": "6", "is_holiday": True},
            {"id": dropped.pk, "version": 1, "excluded": True},
        ],
    )
    before = _snapshot(keep.pk)
    response = client_for(jefe).get(f"{BASE}/depuration/export/", {"period": "2026-09"})
    assert response.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in ws[1]] == OVERTIME_EXPORT_HEADERS  # mismo formato
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    assert len(data) == 1  # la fila excluida no sale
    assert data[0][1] == "0102030405" and data[0][7] == 6  # valor depurado guardado
    assert data[0][9:12] == ["X", None, "X"]  # sábado y feriado con X
    assert _snapshot(keep.pk) == before  # exportar no modifica nada


def test_history_endpoint_is_read_only(client_for, admin, supervisor, jefe):
    _sep_record(admin, _collab(supervisor))
    DepurationService.create(access(jefe), SEP)
    row = DepuratedOvertime.objects.get()
    DepurationService.save_rows(
        access(jefe),
        SEP,
        [{"id": row.pk, "version": 1, "overtime_hours": "3", "reason": "Revisión"}],
    )
    client = client_for(jefe)
    history = client.get(f"{BASE}/depuration/history/", {"period": "2026-09"}).json()
    assert (
        history[0]["old_value"] == "10"
        and history[0]["new_value"] == "3"
        and history[0]["actor"] == "Jefe"
    )
    assert client.post(f"{BASE}/depuration/history/", {}, format="json").status_code == 405
    assert client.delete(f"{BASE}/depuration/history/").status_code == 405
