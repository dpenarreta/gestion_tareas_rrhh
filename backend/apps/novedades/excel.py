"""Importación y exportación de Excel del módulo (openpyxl, mismo motor que
`apps.tasks.task_import`).

La importación es en dos pasos: `ImportService.preview` valida TODO el
archivo sin escribir nada y guarda el resultado en un `ImportBatch`;
`ImportService.confirm` aplica solo las filas válidas con la estrategia
elegida (omitir o actualizar existentes). Confirmar dos veces el mismo
lote no vuelve a aplicar nada, y reimportar el mismo archivo no duplica
registros porque cada fila se identifica por (colaborador, fecha).
"""

import hashlib
import io
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal

import openpyxl
from django.db import IntegrityError, transaction
from django.utils import timezone
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.configuration.services import get_holiday_set
from apps.core.audit import record_audit_event
from apps.tasks.task_import import parse_date

from . import access as perms
from .access import NovedadesAccess
from .models import (
    Collaborator,
    CollaboratorAssignment,
    ImportBatch,
    OvertimeRecord,
    SupervisorPeriod,
    WorkedDay,
)
from .periods import WEEKDAY_NAMES_ES, MonthKey, PeriodKey, is_weekend
from .services import (
    ALERT_NEW_COLLABORATOR,
    EventService,
    NotificationService,
    NovedadesError,
    OvertimeMonthService,
    OvertimeService,
    PeriodService,
    RecordDeadlineService,
    WorkedDayService,
    _fmt,
    _forbidden,
    business_stamp,
    normalize_cedula,
    parse_hours,
    user_display,
)

WORKED_DAYS_FIXED_HEADERS = ["Supervisor", "Cargo", "Cédula", "Empleado"]
# Columna retirada del formato (2026-10-12): si un archivo antiguo la trae,
# la importación la ignora sin error.
LEGACY_IGNORED_HEADERS = {"CORTE ANTERIOR"}
OVERTIME_HEADERS = [
    "FECHA",
    "CEDULA",
    "NOMBRE DEL COLABORADOR",
    "REGIONAL",
    "ÁREA",
    "CARGO",
    "# HORAS EXTRAS",
    "# HORAS SABADO",
    "SABADO",
    "DOMINGO",
    "FERIADO",
    "TRABAJO REALIZADO",
    "OBSERVACIONES",
]
# Exportación del consolidado: SUPERVISOR después de CARGO (7.ª columna).
OVERTIME_EXPORT_HEADERS = OVERTIME_HEADERS[:6] + ["SUPERVISOR"] + OVERTIME_HEADERS[6:]

WEEKEND_FILL = PatternFill("solid", fgColor="FFF59D")
HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
MAX_IMPORT_ROWS = 5000


def _norm(text) -> str:
    """Encabezado comparable: sin tildes, mayúsculas, espacios simples."""
    value = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return " ".join(value.upper().split())


def _overtime_sheet(title: str, rows) -> bytes:
    """Hoja con el formato del reporte de Horas extras (14 columnas, SUPERVISOR
    séptima, fecha dd/mm/aaaa, cédula como texto, X o vacío en SABADO/DOMINGO/
    FERIADO). La usan el reporte original y el de Horas depuradas."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(OVERTIME_EXPORT_HEADERS)
    for (
        day,
        cedula,
        name,
        regional,
        area,
        cargo,
        supervisor,
        extra,
        saturday,
        sat,
        sun,
        hol,
        work,
        obs,
    ) in rows:
        ws.append(
            [
                day,
                str(cedula),
                _safe_text(name),
                _safe_text(regional),
                _safe_text(area),
                _safe_text(cargo),
                _safe_text(supervisor),
                float(extra),
                float(saturday),
                _mark(sat),
                _mark(sun),
                _mark(hol),
                _safe_text(work),
                _safe_text(obs),
            ]
        )
        ws.cell(row=ws.max_row, column=1).number_format = "DD/MM/YYYY"
        _text_cell(ws, ws.max_row, 2)
    _style_header(ws, len(OVERTIME_EXPORT_HEADERS))
    for idx, width in enumerate([12, 14, 32, 16, 18, 20, 26, 10, 10, 8, 8, 8, 36, 36], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    return _sheet_bytes(wb)


def _mark(value: bool | None) -> str | None:
    """SABADO/DOMINGO/FERIADO en el REPORTE exportado: "X" si es verdadero y
    celda vacía si es falso o no hay dato (formato pedido, 2026-10-13). Solo
    presentación: en la base siguen siendo booleanos, y la importación ya lee
    "X" como Sí y la celda vacía como No, así que el reporte se puede volver
    a importar. Solo `True` produce "X": nada ambiguo se toma como Sí."""
    return "X" if value is True else None


def _bool(value) -> bool | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return False
    if isinstance(value, bool):
        return value
    text = _norm(value)
    if text in {"SI", "S", "X", "1", "TRUE", "VERDADERO", "YES"}:
        return True
    if text in {"NO", "N", "0", "FALSE", "FALSO"}:
        return False
    return None


def _header_date(value) -> date | None:
    """Fecha de una celda/encabezado: date real, serial de Excel o texto
    (ISO o dd/mm/aaaa) — mismo parser que la importación de Tareas."""
    return parse_date(value)


def _sheet_bytes(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _safe_text(value) -> str:
    """Texto libre exportado a Excel: un valor que empieza con `=`, `+`,
    `-` o `@` se ejecutaría como fórmula al abrir el archivo (inyección
    CSV/Excel). Se antepone un apóstrofo, que Excel muestra como texto."""
    text = str(value or "")
    return f"'{text}" if text.startswith(_FORMULA_PREFIXES) else text


def _text_cell(ws, row: int, col: int):
    """Cédulas como texto: conserva ceros iniciales y evita notación científica."""
    ws.cell(row=row, column=col).number_format = "@"


def _style_header(ws, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"


class ExportService:
    @staticmethod
    def overtime(acc: NovedadesAccess, qs, *, filters: dict | None = None, context=None) -> bytes:
        """Exporta TODO lo que cumple los filtros dentro del ámbito del
        usuario (no la página visible): `qs` ya viene acotado por
        `OvertimeService.queryset_for`. Queda auditado."""
        if not acc.can(perms.EXPORTAR):
            raise _forbidden("No tiene habilitada la exportación a Excel.")
        EventService.log(
            actor=acc.user,
            action="export.overtime",
            supervisor=acc.user if acc.is_supervisor else None,
            target_type="export",
            target_label="Horas extras",
            new_value=f"{qs.count()} registros",
            reason=", ".join(f"{k}={v}" for k, v in (filters or {}).items() if v)[:500],
            context=context,
        )
        return _overtime_sheet(
            "Horas extras",
            (
                (
                    r.date,
                    r.collaborator.cedula,
                    r.collaborator.full_name,
                    r.collaborator.regional,
                    r.collaborator.area,
                    r.collaborator.cargo,
                    user_display(r.supervisor),
                    r.overtime_hours,
                    r.saturday_hours,
                    r.is_saturday,
                    r.is_sunday,
                    r.is_holiday,
                    r.work_done,
                    r.observations,
                )
                for r in qs.iterator(chunk_size=1000)
            ),
        )

    @staticmethod
    def depurated(acc: NovedadesAccess, depuration, qs, *, context=None) -> bytes:
        """Exporta SOLO la base de Horas depuradas (valores depurados ya
        guardados en el servidor; sin filas excluidas), con el mismo formato
        que el reporte de Horas extras. No escribe en ninguna base."""
        EventService.log(
            actor=acc.user,
            action="export.depuration",
            target_type="overtime_depuration",
            target_id=depuration.pk,
            target_label=f"Horas depuradas {depuration.month:02d}/{depuration.year}",
            new_value=f"{qs.count()} registros",
            context=context,
        )
        return _overtime_sheet(
            "Horas depuradas",
            (
                (
                    r.date,
                    r.cedula,
                    r.full_name,
                    r.regional,
                    r.area,
                    r.cargo,
                    r.supervisor_name,
                    r.overtime_hours,
                    r.saturday_hours,
                    r.is_saturday,
                    r.is_sunday,
                    r.is_holiday,
                    r.work_done,
                    r.observations,
                )
                for r in qs.iterator(chunk_size=1000)
            ),
        )

    @staticmethod
    def overtime_template() -> bytes:
        """Plantilla de importación de Horas extras: los 13 encabezados exactos."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Horas extras"
        ws.append(OVERTIME_HEADERS)
        _style_header(ws, len(OVERTIME_HEADERS))
        for idx, width in enumerate([12, 14, 32, 16, 18, 20, 10, 10, 8, 8, 8, 36, 36], start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for r in range(2, 502):
            _text_cell(ws, r, 2)
            ws.cell(row=r, column=1).number_format = "DD/MM/YYYY"
        return _sheet_bytes(wb)

    @staticmethod
    def worked_days(acc: NovedadesAccess, matrix: dict, *, blank: bool = False) -> bytes:
        """Exportación de la matriz; con `blank=True` es la plantilla de
        importación (mismos colaboradores y fechas, celdas vacías)."""
        if not blank and not acc.can(perms.EXPORTAR):
            raise _forbidden("No tiene habilitada la exportación a Excel.")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Días laborados"
        columns = matrix["columns"]
        dates = [date.fromisoformat(c["date"]) for c in columns]
        # Encabezados de fecha como TEXTO "23/09/2026": se ven igual en
        # cualquier configuración regional de Excel y el importador los lee
        # con el mismo parser que acepta fechas reales de Excel.
        ws.append(WORKED_DAYS_FIXED_HEADERS + [f"{d:%d/%m/%Y}" for d in dates])
        fixed = len(WORKED_DAYS_FIXED_HEADERS)
        for idx in range(fixed + 1, fixed + len(dates) + 1):
            ws.cell(row=1, column=idx).number_format = "@"
        supervisor_name = matrix["supervisor"]["name"]
        for row in matrix["rows"]:
            values = []
            for c in columns:
                cell = row["cells"].get(c["date"])
                values.append(float(cell["hours"]) if cell and cell["hours"] is not None else None)
            ws.append(
                [
                    _safe_text(supervisor_name),
                    _safe_text(row["cargo"]),
                    str(row["cedula"]),
                    _safe_text(row["full_name"]),
                ]
                + (values if not blank else [None] * len(values))
            )
            _text_cell(ws, ws.max_row, 3)
        _style_header(ws, fixed + len(dates))
        for idx, d in enumerate(dates, start=fixed + 1):
            ws.column_dimensions[get_column_letter(idx)].width = 11
            if is_weekend(d):
                for r in range(1, ws.max_row + 1):
                    ws.cell(row=r, column=idx).fill = WEEKEND_FILL
        for idx, width in enumerate([26, 20, 14, 32], start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = ws.cell(row=2, column=fixed + 1)
        if blank:
            ExportService._worked_days_instructions(wb, matrix, dates)
        return _sheet_bytes(wb)

    @staticmethod
    def _worked_days_instructions(wb, matrix: dict, dates: list[date]) -> None:
        period = matrix["period"]
        ws = wb.create_sheet("Instrucciones")
        lines = [
            f"Plantilla de Días laborados — {period['label']}",
            f"Período: {dates[0]:%d/%m/%Y} al {dates[-1]:%d/%m/%Y} (corte del 23 al 22).",
            "",
            "Columnas fijas (no cambie sus nombres): "
            + " | ".join(WORKED_DAYS_FIXED_HEADERS)
            + ".",
            "• Cédula: identifica a la persona (debe existir en la Base de colaboradores del equipo).",
            "• Empleado: nombre completo; debe corresponder a la cédula (se avisa si no coincide).",
            "• Supervisor y Cargo son informativos: se toman de la base, no del archivo.",
            "",
            "Columnas de fecha: una por cada día del período, con encabezado dd/mm/aaaa.",
            "• Escriba las HORAS laboradas de ese día: número de 0 a 24, hasta 2 decimales (ej. 8 o 7.5).",
            "• Celda vacía = no reportado. 0 = reportó cero horas (no son lo mismo).",
            "• Las columnas en amarillo (sábado y domingo) deben quedar vacías: no admiten registros.",
            "• No agregue fechas fuera del período ni repita una fecha.",
            "",
            "Si el período ya está cerrado para el equipo, la importación no se aplica.",
        ]
        for line in lines:
            ws.append([line])
        ws["A1"].font = Font(bold=True)
        ws.column_dimensions["A"].width = 110

    @staticmethod
    def import_errors(batch: ImportBatch) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Errores"
        ws.append(["Fila", "Columna", "Error"])
        for err in batch.errors:
            ws.append([err.get("row"), err.get("column", ""), err.get("message")])
        _style_header(ws, 3)
        ws.column_dimensions["C"].width = 90
        return _sheet_bytes(wb)


class ImportService:
    @staticmethod
    def _check(acc: NovedadesAccess):
        if not acc.can(perms.IMPORTAR):
            raise _forbidden("No tiene habilitada la importación de Excel.")
        if not acc.can(perms.REGISTRAR):
            raise _forbidden("No tiene habilitado el registro de novedades.")
        if acc.is_assistant:
            raise _forbidden("El Asistente de Operaciones no modifica registros originales.")

    @staticmethod
    def _load(file) -> list[tuple]:
        try:
            wb = openpyxl.load_workbook(file, data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise NovedadesError(
                "No se pudo leer el archivo: debe ser un Excel .xlsx válido."
            ) from exc
        ws = wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise NovedadesError("El archivo está vacío.")
        if len(rows) - 1 > MAX_IMPORT_ROWS:
            raise NovedadesError(f"El archivo supera el máximo de {MAX_IMPORT_ROWS} filas.")
        return rows

    @staticmethod
    def _collaborator_for(acc, cedula, by_cedula, errors, row_num, column="Cédula"):
        collaborator = by_cedula.get(cedula)
        if collaborator is None:
            errors.append(
                {
                    "row": row_num,
                    "column": column,
                    "message": f"La cédula {cedula} no existe en la base de colaboradores. Cárguela primero en Base de colaboradores (puede importarla desde ahí).",
                }
            )
            return None
        if not acc.can_write_for_supervisor(collaborator.current_supervisor_id):
            errors.append(
                {
                    "row": row_num,
                    "column": column,
                    "message": f"La cédula {cedula} no pertenece a su equipo.",
                }
            )
            return None
        return collaborator

    @staticmethod
    def _open_supervisor_ids(key: PeriodKey, supervisor_ids: set[int]) -> set[int]:
        for sid in supervisor_ids:
            from django.contrib.auth import get_user_model

            PeriodService.ensure_current(get_user_model().objects.get(pk=sid))
        return set(
            SupervisorPeriod.objects.filter(
                period__year=key.year,
                period__month=key.month,
                supervisor_id__in=supervisor_ids,
                status=SupervisorPeriod.Status.OPEN,
            ).values_list("supervisor_id", flat=True)
        )

    @staticmethod
    def preview(
        acc: NovedadesAccess, *, kind: str, key: PeriodKey, file, file_name: str
    ) -> ImportBatch:
        ImportService._check(acc)
        if kind == ImportBatch.Kind.OVERTIME:
            # Horas extras se importa por MES CALENDARIO (1 al último día).
            key = MonthKey(key.year, key.month)
        content = file.read()
        file_hash = hashlib.sha256(content).hexdigest()
        rows = ImportService._load(io.BytesIO(content))
        warnings = []
        if kind == ImportBatch.Kind.WORKED_DAYS:
            parsed, errors, warnings = ImportService._parse_worked_days(acc, key, rows)
        elif kind == ImportBatch.Kind.OVERTIME:
            parsed, errors = ImportService._parse_overtime(acc, key, rows)
            if OvertimeMonthService.current_month() < key:
                warnings.append(
                    {
                        "row": 1,
                        "column": "FECHA",
                        "message": f"{key.label} todavía no empieza: la confirmación será rechazada.",
                    }
                )
            closed = OvertimeMonthService.closed_supervisor_ids(
                key, {p["supervisor_id"] for p in parsed}
            )
            if closed:
                warnings.append(
                    {
                        "row": 1,
                        "column": "FECHA",
                        "message": f"{key.label} de Horas extras está cerrado para "
                        f"{len(closed)} equipo(s) del archivo: la confirmación será rechazada.",
                    }
                )
        else:
            raise NovedadesError("Tipo de importación inválido.")
        already = ImportBatch.objects.filter(
            kind=kind,
            period__year=key.year,
            period__month=key.month,
            file_hash=file_hash,
            status=ImportBatch.Status.CONFIRMED,
        ).exists()
        error_rows = {e["row"] for e in errors}
        summary = {
            "total_rows": len(rows) - 1,
            "valid_rows": len({p["row"] for p in parsed} - error_rows),
            "invalid_rows": len(error_rows),
            "existing_items": sum(1 for p in parsed if p.get("existing")),
            "new_items": sum(1 for p in parsed if not p.get("existing")),
            "already_imported": already,
            "warnings": warnings,
        }
        # Filas con algún error se descartan completas (nunca a medias).
        parsed = [p for p in parsed if p["row"] not in error_rows]
        return ImportBatch.objects.create(
            kind=kind,
            # Mismo año/mes; para Horas extras se interpreta como mes calendario.
            period=PeriodService.get_period(PeriodKey(key.year, key.month)),
            uploaded_by=acc.user,
            file_name=file_name[:255],
            file_hash=file_hash,
            rows=parsed,
            errors=errors,
            summary=summary,
        )

    @staticmethod
    def _parse_worked_days(acc, key: PeriodKey, rows):
        deadline_holidays = get_holiday_set()
        header = list(rows[0])
        errors, parsed = [], []
        # Las columnas fijas se ubican por su encabezado: así un archivo del
        # formato anterior (con "Corte anterior") sigue siendo válido y esa
        # columna simplemente se ignora.
        positions = {
            _norm(h): i for i, h in enumerate(header) if h is not None and not _header_date(h)
        }
        missing = [h for h in WORKED_DAYS_FIXED_HEADERS if _norm(h) not in positions]
        if missing:
            raise NovedadesError(
                "Encabezados inválidos. Faltan: "
                + " | ".join(missing)
                + ". Columnas esperadas: "
                + " | ".join(WORKED_DAYS_FIXED_HEADERS)
                + " | una columna por fecha del período."
            )
        cedula_idx = positions[_norm("Cédula")]
        name_idx = positions[_norm("Empleado")]
        fixed_idx = {positions[_norm(h)] for h in WORKED_DAYS_FIXED_HEADERS} | {
            i for name, i in positions.items() if name in LEGACY_IGNORED_HEADERS
        }
        date_cols: list[tuple[int, date]] = []
        for idx, raw in enumerate(header):
            if idx in fixed_idx or raw is None or str(raw).strip() == "":
                continue
            d = _header_date(raw)
            if d is None:
                errors.append(
                    {"row": 1, "column": str(raw), "message": "Encabezado de fecha inválido."}
                )
                continue
            if not key.contains(d):
                errors.append(
                    {
                        "row": 1,
                        "column": f"{d:%d/%m/%Y}",
                        "message": f"La fecha no pertenece al período {key.label} "
                        f"({key.start_date:%d/%m/%Y} – {key.end_date:%d/%m/%Y}).",
                    }
                )
                continue
            if any(d == other for _, other in date_cols):
                errors.append(
                    {
                        "row": 1,
                        "column": f"{d:%d/%m/%Y}",
                        "message": "Fecha repetida en los encabezados.",
                    }
                )
                continue
            date_cols.append((idx, d))
        if errors:
            return [], errors, []
        present = {d for _, d in date_cols}
        warnings = []
        missing_days = [d for d in key.dates() if d not in present and not is_weekend(d)]
        if missing_days:
            warnings.append(
                {
                    "row": 1,
                    "column": "Fechas",
                    "message": f"El archivo no trae {len(missing_days)} día(s) hábil(es) del período "
                    f"(p. ej. {missing_days[0]:%d/%m/%Y}); esas celdas no se modifican.",
                }
            )

        cedulas = {
            normalize_import_cedula(r[cedula_idx])
            for r in rows[1:]
            if len(r) > cedula_idx and r[cedula_idx] not in (None, "")
        }
        by_cedula = {c.cedula: c for c in Collaborator.objects.filter(cedula__in=cedulas)}
        existing = {
            (wd.collaborator_id, wd.date): wd
            for wd in WorkedDay.objects.filter(
                collaborator__cedula__in=cedulas, date__gte=key.start_date, date__lte=key.end_date
            )
        }
        seen = set()
        for offset, row in enumerate(rows[1:], start=2):
            if all(v in (None, "") for v in row):
                continue
            cedula = normalize_import_cedula(row[cedula_idx] if len(row) > cedula_idx else "")
            if not cedula:
                errors.append({"row": offset, "column": "Cédula", "message": "Cédula vacía."})
                continue
            if cedula in seen:
                errors.append(
                    {
                        "row": offset,
                        "column": "Cédula",
                        "message": f"Cédula {cedula} repetida en el archivo.",
                    }
                )
                continue
            seen.add(cedula)
            collaborator = ImportService._collaborator_for(acc, cedula, by_cedula, errors, offset)
            if collaborator is None:
                continue
            file_name = " ".join(str(row[name_idx] or "").split()) if name_idx < len(row) else ""
            name_check = ImportService._name_matches(file_name, collaborator.full_name)
            if name_check == "different":
                errors.append(
                    {
                        "row": offset,
                        "column": "Empleado",
                        "message": f"La cédula {cedula} corresponde a {collaborator.full_name} en la base, "
                        f"pero el archivo dice «{file_name}». Revise la cédula.",
                    }
                )
                continue
            if name_check == "partial":
                warnings.append(
                    {
                        "row": offset,
                        "column": "Empleado",
                        "message": f"«{file_name}» no coincide exactamente con «{collaborator.full_name}» "
                        "(cédula correcta en la base): se usará el nombre de la base.",
                    }
                )
            for idx, d in date_cols:
                raw = row[idx] if idx < len(row) else None
                if raw is None or (isinstance(raw, str) and raw.strip() == ""):
                    continue
                column = f"{d:%d/%m/%Y}"
                if is_weekend(d):
                    errors.append(
                        {
                            "row": offset,
                            "column": column,
                            "message": f"{WEEKDAY_NAMES_ES[d.weekday()]}: las columnas de sábado y domingo no "
                            "admiten valores en Días laborados.",
                        }
                    )
                    continue
                try:
                    # Mismo plazo por fecha y prohibición de fechas futuras que
                    # la captura: se avisa en la vista previa, no al confirmar.
                    RecordDeadlineService.check(
                        acc, RecordDeadlineService.WORKED_DAYS, d, holidays=deadline_holidays
                    )
                    value = parse_hours(raw, field_label="Horas")
                except NovedadesError as exc:
                    errors.append({"row": offset, "column": column, "message": exc.message})
                    continue
                if collaborator.status != Collaborator.Status.ACTIVE:
                    errors.append(
                        {
                            "row": offset,
                            "column": column,
                            "message": f"{collaborator.full_name} está {collaborator.get_status_display().lower()}.",
                        }
                    )
                    continue
                current = existing.get((collaborator.pk, d))
                parsed.append(
                    {
                        "row": offset,
                        "collaborator_id": collaborator.pk,
                        "supervisor_id": collaborator.current_supervisor_id,
                        "cedula": cedula,
                        "name": collaborator.full_name,
                        "date": d.isoformat(),
                        "hours": _fmt(value),
                        "existing": current is not None and current.hours is not None,
                        "existing_hours": (
                            _fmt(current.hours) if current and current.hours is not None else None
                        ),
                    }
                )
        return parsed, errors, warnings

    @staticmethod
    def _name_matches(file_name: str, official: str) -> str:
        """`same` / `partial` (mismas palabras en otro orden u omitidas) /
        `different` (ninguna palabra en común: probable cédula equivocada).
        Un nombre vacío en el archivo no se objeta: manda la cédula."""
        if not file_name:
            return "same"
        a, b = _norm(file_name), _norm(official)
        if a == b:
            return "same"
        words_a = {w for w in a.split() if len(w) > 2}
        words_b = {w for w in b.split() if len(w) > 2}
        return "partial" if words_a & words_b else "different"

    @staticmethod
    def _parse_overtime(acc, key: PeriodKey, rows):
        deadline_holidays = get_holiday_set()
        header = [_norm(h) for h in rows[0][: len(OVERTIME_HEADERS)]]
        if header != [_norm(h) for h in OVERTIME_HEADERS]:
            raise NovedadesError(
                "Encabezados inválidos. Deben ser exactamente: " + " | ".join(OVERTIME_HEADERS)
            )
        errors, parsed = [], []
        cedulas = {
            normalize_import_cedula(r[1]) for r in rows[1:] if len(r) > 1 and r[1] not in (None, "")
        }
        by_cedula = {c.cedula: c for c in Collaborator.objects.filter(cedula__in=cedulas)}
        existing = {
            (r.collaborator_id, r.date): r
            for r in OvertimeRecord.objects.filter(
                collaborator__cedula__in=cedulas,
                date__gte=key.start_date,
                date__lte=key.end_date,
                deleted_at__isnull=True,
            )
        }
        seen = set()
        for offset, row in enumerate(rows[1:], start=2):
            row = list(row) + [None] * (len(OVERTIME_HEADERS) - len(row))
            if all(v in (None, "") for v in row):
                continue
            d = _header_date(row[0])
            if d is None:
                errors.append({"row": offset, "column": "FECHA", "message": "Fecha inválida."})
                continue
            if not key.contains(d):
                errors.append(
                    {
                        "row": offset,
                        "column": "FECHA",
                        "message": f"{d:%d/%m/%Y} no pertenece a {key.label}.",
                    }
                )
                continue
            cedula = normalize_import_cedula(row[1])
            if not cedula:
                errors.append({"row": offset, "column": "CEDULA", "message": "Cédula vacía."})
                continue
            collaborator = ImportService._collaborator_for(
                acc, cedula, by_cedula, errors, offset, "CEDULA"
            )
            if collaborator is None:
                continue
            if collaborator.status != Collaborator.Status.ACTIVE:
                errors.append(
                    {
                        "row": offset,
                        "column": "CEDULA",
                        "message": f"{collaborator.full_name} está {collaborator.get_status_display().lower()}.",
                    }
                )
                continue
            if (cedula, d) in seen:
                errors.append(
                    {
                        "row": offset,
                        "column": "FECHA",
                        "message": f"Registro duplicado en el archivo ({cedula}, {d:%d/%m/%Y}).",
                    }
                )
                continue
            seen.add((cedula, d))
            try:
                RecordDeadlineService.check(
                    acc, RecordDeadlineService.OVERTIME, d, holidays=deadline_holidays
                )
            except NovedadesError as exc:
                errors.append({"row": offset, "column": "FECHA", "message": exc.message})
                continue
            try:
                overtime = parse_hours(row[6], field_label="# HORAS EXTRAS") or Decimal("0")
                saturday = parse_hours(row[7], field_label="# HORAS SABADO") or Decimal("0")
            except NovedadesError as exc:
                errors.append({"row": offset, "column": "HORAS", "message": exc.message})
                continue
            if overtime + saturday <= 0:
                errors.append(
                    {"row": offset, "column": "# HORAS EXTRAS", "message": "Sin horas registradas."}
                )
                continue
            flags = {}
            bad_flag = False
            for col, name in ((8, "is_saturday"), (9, "is_sunday"), (10, "is_holiday")):
                value = _bool(row[col])
                if value is None:
                    errors.append(
                        {
                            "row": offset,
                            "column": OVERTIME_HEADERS[col],
                            "message": "Use SI o NO (o deje vacío).",
                        }
                    )
                    bad_flag = True
                flags[name] = value
            if bad_flag:
                continue
            current = existing.get((collaborator.pk, d))
            parsed.append(
                {
                    "row": offset,
                    "collaborator_id": collaborator.pk,
                    "supervisor_id": collaborator.current_supervisor_id,
                    "cedula": cedula,
                    "name": collaborator.full_name,
                    "date": d.isoformat(),
                    "overtime_hours": _fmt(overtime),
                    "saturday_hours": _fmt(saturday),
                    **flags,
                    "work_done": str(row[11] or "").strip()[:500],
                    "observations": str(row[12] or "").strip()[:1000],
                    "existing": current is not None,
                    "existing_id": current.pk if current else None,
                }
            )
        return parsed, errors

    @staticmethod
    def confirm(acc: NovedadesAccess, batch_id, strategy: str, *, context=None) -> ImportBatch:
        batch = (
            ImportBatch.objects.select_related("period")
            .filter(pk=batch_id, uploaded_by=acc.user)
            .first()
        )
        if batch is None:
            raise NovedadesError("Importación no encontrada.", code="not_found", status=404)
        if batch.status == ImportBatch.Status.CONFIRMED:
            return batch  # idempotente
        if batch.status == ImportBatch.Status.CANCELLED:
            raise NovedadesError("La importación fue cancelada.", code="cancelled", status=409)
        if strategy not in ("skip", "update"):
            raise NovedadesError(
                "Estrategia inválida: use 'skip' (omitir existentes) o 'update' (actualizar)."
            )
        if not ImportService._claim(batch.pk):
            confirmed = ImportBatch.objects.get(pk=batch.pk)
            if confirmed.status == ImportBatch.Status.CONFIRMED:
                return confirmed  # otra solicitud terminó de confirmarlo
            raise NovedadesError(
                "Esta importación ya se está confirmando; espere el resultado.",
                code="import_in_progress",
                status=409,
            )
        try:
            return ImportService._confirm_claimed(acc, batch, strategy, context=context)
        finally:
            ImportService._release(batch.pk)

    CLAIM_MINUTES = 10

    @staticmethod
    def _claim(batch_id) -> bool:
        """Marca el lote "en confirmación" con la fila bloqueada. Devuelve False
        si ya está confirmado o si otra solicitud lo está confirmando. Una
        marca de más de 10 minutos se considera abandonada (proceso caído)."""
        with transaction.atomic():
            locked = ImportBatch.objects.select_for_update().get(pk=batch_id)
            if locked.status != ImportBatch.Status.PREVIEW:
                return False
            since = (locked.summary or {}).get("confirming_since")
            now = timezone.now()
            if since:
                try:
                    started = datetime.fromisoformat(since)
                except ValueError:
                    started = None
                if started and now - started < timedelta(minutes=ImportService.CLAIM_MINUTES):
                    return False
            locked.summary = {**(locked.summary or {}), "confirming_since": now.isoformat()}
            locked.save(update_fields=["summary"])
        return True

    @staticmethod
    def _release(batch_id) -> None:
        batch = ImportBatch.objects.filter(pk=batch_id).first()
        if batch and "confirming_since" in (batch.summary or {}):
            batch.summary = {k: v for k, v in batch.summary.items() if k != "confirming_since"}
            batch.save(update_fields=["summary"])

    @staticmethod
    def _confirm_claimed(acc: NovedadesAccess, batch, strategy: str, *, context=None):
        if batch.kind == ImportBatch.Kind.COLLABORATORS:
            # La base de colaboradores CREA personas: no pasa por la exigencia
            # de existencia previa de las importaciones de novedades.
            return CollaboratorImportService.confirm(acc, batch, strategy, context=context)
        ImportService._check(acc)
        key = batch.period.key
        items = [i for i in batch.rows if strategy == "update" or not i.get("existing")]
        supervisor_ids = {i["supervisor_id"] for i in items}
        if batch.kind == ImportBatch.Kind.OVERTIME:
            # Mes calendario; bloquea solo el cierre manual del mes (no hay
            # cierre automático por plazo vencido).
            key = MonthKey(key.year, key.month)
            open_ids = supervisor_ids - OvertimeMonthService.closed_supervisor_ids(
                key, supervisor_ids
            )
        else:
            open_ids = ImportService._open_supervisor_ids(key, supervisor_ids)
        closed = [i for i in items if i["supervisor_id"] not in open_ids]
        if closed:
            raise NovedadesError(
                f"{key.label} está cerrado para alguno de los supervisores del archivo; no se importó nada.",
                code="period_closed",
                status=409,
            )
        applied = skipped = 0
        from django.contrib.auth import get_user_model

        users = {u.pk: u for u in get_user_model().objects.filter(pk__in=open_ids)}
        with transaction.atomic():
            if batch.kind == ImportBatch.Kind.WORKED_DAYS:
                by_supervisor: dict[int, list] = {}
                for item in items:
                    by_supervisor.setdefault(item["supervisor_id"], []).append(
                        {
                            "collaborator_id": item["collaborator_id"],
                            "date": date.fromisoformat(item["date"]),
                            "hours": item["hours"],
                        }
                    )
                for sid, cells in by_supervisor.items():
                    result = WorkedDayService.save_cells(
                        acc,
                        key,
                        users[sid],
                        cells,
                        source_file=batch.file_name,
                        context=context,
                        notify=False,
                    )
                    applied += result["saved"]
            elif items:
                # Misma validación y misma escritura que la captura masiva
                # (`OvertimeService.bulk_save`): la importación no tiene un
                # camino propio que pueda divergir.
                day_of = {i["row"]: date.fromisoformat(i["date"]) for i in items}
                current = {
                    (r.collaborator_id, r.date): r
                    for r in OvertimeRecord.objects.filter(
                        collaborator_id__in={i["collaborator_id"] for i in items},
                        date__in=set(day_of.values()),
                        deleted_at__isnull=True,
                    )
                }
                rows = []
                for item in items:
                    existing = current.get((item["collaborator_id"], day_of[item["row"]]))
                    if existing is not None and strategy == "skip":
                        # Se registró después de la vista previa: "omitir
                        # existentes" no lo sobrescribe.
                        skipped += 1
                        continue
                    rows.append(
                        {
                            "client_key": str(item["row"]),
                            "id": existing.pk if existing else None,
                            "version": existing.version if existing else None,
                            "collaborator_id": item["collaborator_id"],
                            "date": day_of[item["row"]],
                            **{
                                f: item[f]
                                for f in (
                                    "overtime_hours",
                                    "saturday_hours",
                                    "is_saturday",
                                    "is_sunday",
                                    "is_holiday",
                                    "work_done",
                                    "observations",
                                )
                            },
                        }
                    )
                for start in range(0, len(rows), OvertimeService.BULK_LIMIT):
                    result = OvertimeService.bulk_save(
                        acc,
                        key,
                        rows[start : start + OvertimeService.BULK_LIMIT],
                        source_file=batch.file_name,
                        context=context,
                        notify=False,
                    )
                    applied += result["summary"]["created"] + result["summary"]["updated"]
                    skipped += result["summary"]["unchanged"]
                    for r in result["rows"]:
                        if r["errors"]:
                            batch.errors.append(
                                {
                                    "row": int(r["client_key"]),
                                    "column": "",
                                    "message": " ".join(r["errors"].values()),
                                }
                            )
            skipped += len(batch.rows) - len(items)
            batch.status = ImportBatch.Status.CONFIRMED
            batch.strategy = strategy
            batch.summary = {**batch.summary, "applied": applied, "skipped": skipped}
            batch.save(update_fields=["status", "strategy", "summary", "errors", "updated_at"])
            EventService.log(
                actor=acc.user,
                action="import.confirmed",
                period=batch.period,
                supervisor=acc.user if acc.is_supervisor else None,
                target_type="import",
                target_id=batch.pk,
                target_label=batch.file_name,
                new_value=f"{applied} aplicados, {skipped} omitidos, {len(batch.errors)} errores",
                source_file=batch.file_name,
                context=context,
            )
            NotificationService.notify_admins(
                f"Importación de {batch.get_kind_display()}: {user_display(acc.user)} importó "
                f"'{batch.file_name}' en {key.label} — {applied} aplicados, {skipped} omitidos, "
                f"{len(batch.errors)} con error.",
                exclude_user=acc.user,
            )
        return batch

    @staticmethod
    def cancel(acc: NovedadesAccess, batch_id) -> None:
        ImportBatch.objects.filter(
            pk=batch_id, uploaded_by=acc.user, status=ImportBatch.Status.PREVIEW
        ).update(status=ImportBatch.Status.CANCELLED)


# ── Importación de la base de colaboradores ───────────────────────────────────

COLLABORATOR_TEMPLATE_HEADERS = ["CEDULA", "NOMBRES COMPLETOS", "REGIONAL", "ÁREA", "CARGO"]
# Sinónimos aceptados por columna (comparados sin tildes ni mayúsculas), para
# poder cargar la base desde las planillas que la operación ya usa —incluida
# la de Días laborados (columna "Empleado")—. Cualquier otra columna se ignora.
COLLABORATOR_HEADER_ALIASES = {
    "cedula": {"CEDULA", "CEDULA DE IDENTIDAD", "IDENTIFICACION", "NUMERO DE CEDULA", "CI"},
    "full_name": {
        "NOMBRES COMPLETOS",
        "NOMBRE COMPLETO",
        "NOMBRE DEL COLABORADOR",
        "APELLIDOS Y NOMBRES",
        "NOMBRES Y APELLIDOS",
        "NOMBRES",
        "NOMBRE",
        "EMPLEADO",
        "COLABORADOR",
    },
    "regional": {"REGIONAL"},
    "area": {"AREA"},
    "cargo": {"CARGO"},
}
# Campos que una importación puede actualizar en un colaborador existente.
# Nunca: cédula, estado operativo, supervisor asignado, ni nada de usuarios,
# roles o permisos (el colaborador operativo no es un usuario de NEXO).
COLLABORATOR_UPDATABLE = ("full_name", "regional", "area", "cargo")
COLLABORATOR_LABELS = {
    "full_name": "Nombre",
    "regional": "Regional",
    "area": "Área",
    "cargo": "Cargo",
}


def normalize_import_cedula(raw) -> str:
    """Cédula como texto. Excel guarda una cédula numérica sin el cero
    inicial (0102030405 → 102030405): una de 9 dígitos se completa a los 10
    de la cédula ecuatoriana. Texto alfanumérico (pasaporte) se respeta."""
    if raw is None:
        return ""
    if isinstance(raw, float) and raw.is_integer():
        raw = int(raw)
    text = normalize_cedula(raw)
    if text.isdigit() and len(text) == 9:
        text = text.zfill(10)
    return text


def cedula_problem(cedula: str) -> str | None:
    if not cedula:
        return "Cédula vacía."
    if not cedula.isalnum():
        return f"Cédula «{cedula}» con caracteres no permitidos."
    if not 6 <= len(cedula) <= 20:
        return f"Cédula «{cedula}» con longitud inválida (6 a 20 caracteres)."
    return None


class CollaboratorImportService:
    @staticmethod
    def _check(acc: NovedadesAccess):
        if not acc.can(perms.IMPORTAR):
            raise _forbidden("No tiene habilitada la importación de Excel.")
        if not acc.can(perms.COLABORADORES) or acc.is_assistant:
            raise _forbidden("No tiene permiso para gestionar la base de colaboradores.")

    @staticmethod
    def _target_supervisor(acc: NovedadesAccess, supervisor_id):
        if acc.is_supervisor and not acc.national_write:
            return acc.user
        if not supervisor_id:
            raise NovedadesError("Seleccione el supervisor al que se asignará la base importada.")
        supervisor = PeriodService.get_supervisor(supervisor_id)
        if not acc.can_write_for_supervisor(supervisor.pk):
            raise _forbidden()
        return supervisor

    @staticmethod
    def template() -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Colaboradores"
        ws.append(COLLABORATOR_TEMPLATE_HEADERS)
        _style_header(ws, len(COLLABORATOR_TEMPLATE_HEADERS))
        for idx, width in enumerate([16, 36, 18, 20, 24], start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        for r in range(2, 1002):
            _text_cell(ws, r, 1)  # cédula como texto: conserva el cero inicial
        return _sheet_bytes(wb)

    @staticmethod
    def _classify(supervisor, rows: list[dict]) -> tuple[list[dict], list[dict]]:
        """Clasifica filas ya leídas como NEW / EXISTING / UNCHANGED o error,
        contra el estado ACTUAL de la base (se usa en la vista previa y otra
        vez al confirmar)."""
        cedulas = {r["cedula"] for r in rows}
        existing = {
            c.cedula: c
            for c in Collaborator.objects.select_related("current_supervisor").filter(
                cedula__in=cedulas
            )
        }
        planned, errors = [], []
        for r in rows:
            current = existing.get(r["cedula"])
            if current is None:
                planned.append({**r, "status": "NEW", "collaborator_id": None, "changes": {}})
                continue
            if current.current_supervisor_id != supervisor.pk:
                errors.append(
                    {
                        "row": r["row"],
                        "column": "CEDULA",
                        "message": f"La cédula {r['cedula']} ya pertenece a la base de "
                        f"{user_display(current.current_supervisor)}: no se reasigna por "
                        "importación (el Administrador puede usar Reasignar).",
                    }
                )
                continue
            changes = {
                f: [getattr(current, f), r[f]]
                for f in COLLABORATOR_UPDATABLE
                if r.get(f) and r[f] != getattr(current, f)
            }
            planned.append(
                {
                    **r,
                    "status": "EXISTING" if changes else "UNCHANGED",
                    "collaborator_id": current.pk,
                    "changes": changes,
                    "current_status": current.status,
                }
            )
        return planned, errors

    @staticmethod
    def preview(acc: NovedadesAccess, *, file, file_name: str, supervisor_id=None) -> ImportBatch:
        CollaboratorImportService._check(acc)
        supervisor = CollaboratorImportService._target_supervisor(acc, supervisor_id)
        content = file.read()
        sheet = ImportService._load(io.BytesIO(content))
        header = [_norm(h) for h in sheet[0]]
        columns = {}
        for field, aliases in COLLABORATOR_HEADER_ALIASES.items():
            for idx, name in enumerate(header):
                if name in aliases:
                    columns[field] = idx
                    break
        missing = [
            label
            for field, label in (("cedula", "CEDULA"), ("full_name", "NOMBRES COMPLETOS"))
            if field not in columns
        ]
        if missing:
            raise NovedadesError(
                "Encabezados inválidos: faltan "
                + " y ".join(missing)
                + ". Columnas reconocidas: "
                + " | ".join(COLLABORATOR_TEMPLATE_HEADERS)
                + " (también se aceptan «Cédula», «Empleado», «Nombre del colaborador», etc.)."
            )
        used = set(columns.values())
        ignored = [str(sheet[0][i]) for i in range(len(header)) if i not in used and header[i]]

        rows, errors, seen = [], [], {}
        duplicates = 0
        data_rows = 0
        for offset, raw_row in enumerate(sheet[1:], start=2):
            raw_row = list(raw_row)
            if all(v in (None, "") for v in raw_row):
                continue
            data_rows += 1

            def cell(field, raw_row=raw_row):
                idx = columns.get(field)
                return raw_row[idx] if idx is not None and idx < len(raw_row) else None

            cedula = normalize_import_cedula(cell("cedula"))
            problem = cedula_problem(cedula)
            if problem:
                errors.append({"row": offset, "column": "CEDULA", "message": problem})
                continue
            full_name = " ".join(str(cell("full_name") or "").split())
            if not full_name:
                errors.append(
                    {"row": offset, "column": "NOMBRES COMPLETOS", "message": "Nombre vacío."}
                )
                continue
            if cedula in seen:
                duplicates += 1
                errors.append(
                    {
                        "row": offset,
                        "column": "CEDULA",
                        "message": f"Cédula {cedula} repetida en el archivo "
                        f"(ya está en la fila {seen[cedula]}).",
                    }
                )
                continue
            seen[cedula] = offset
            values = {
                f: " ".join(str(cell(f) or "").split()) for f in ("regional", "area", "cargo")
            }
            too_long = [
                COLLABORATOR_LABELS[f]
                for f in ("regional", "area", "cargo")
                if len(values[f]) > 100
            ]
            if len(full_name) > 200:
                too_long.append("Nombre")
            if too_long:
                errors.append(
                    {
                        "row": offset,
                        "column": ", ".join(too_long),
                        "message": "Texto demasiado largo.",
                    }
                )
                continue
            rows.append({"row": offset, "cedula": cedula, "full_name": full_name, **values})

        planned, conflicts = CollaboratorImportService._classify(supervisor, rows)
        errors.extend(conflicts)
        summary = {
            "total_rows": data_rows,
            "new_items": sum(1 for p in planned if p["status"] == "NEW"),
            "existing_items": sum(1 for p in planned if p["status"] != "NEW"),
            "to_update": sum(1 for p in planned if p["status"] == "EXISTING"),
            "unchanged": sum(1 for p in planned if p["status"] == "UNCHANGED"),
            "duplicates_in_file": duplicates,
            "invalid_rows": len({e["row"] for e in errors}),
            "valid_rows": len(planned),
            "ignored_columns": ignored,
            "supervisor_id": supervisor.pk,
            "supervisor_name": user_display(supervisor),
        }
        return ImportBatch.objects.create(
            kind=ImportBatch.Kind.COLLABORATORS,
            period=None,
            uploaded_by=acc.user,
            file_name=file_name[:255],
            file_hash=hashlib.sha256(content).hexdigest(),
            rows=planned,
            errors=sorted(errors, key=lambda e: e["row"]),
            summary=summary,
        )

    @staticmethod
    def confirm(
        acc: NovedadesAccess, batch: ImportBatch, strategy: str, *, context=None
    ) -> ImportBatch:
        """Crea los nuevos y, con `strategy="update"`, actualiza los campos
        permitidos de los existentes. Se revalida todo contra la base actual
        (alguien pudo crear la misma cédula entre la vista previa y la
        confirmación); cada fila se guarda en su propia transacción."""
        CollaboratorImportService._check(acc)
        supervisor = CollaboratorImportService._target_supervisor(
            acc, batch.summary.get("supervisor_id")
        )
        rows = [
            {k: r[k] for k in ("row", "cedula", "full_name", "regional", "area", "cargo")}
            for r in batch.rows
        ]
        planned, conflicts = CollaboratorImportService._classify(supervisor, rows)
        current = PeriodService.ensure_current(supervisor)
        created, updated, unchanged, rejected = [], 0, 0, list(conflicts)
        for item in planned:
            try:
                with transaction.atomic():
                    if item["status"] == "NEW":
                        collaborator = Collaborator.objects.create(
                            cedula=item["cedula"],
                            full_name=item["full_name"],
                            regional=item["regional"],
                            area=item["area"],
                            cargo=item["cargo"],
                            current_supervisor=supervisor,
                            created_by=acc.user,
                        )
                        CollaboratorAssignment.objects.create(
                            collaborator=collaborator,
                            supervisor=supervisor,
                            assigned_by=acc.user,
                            reason=f"Alta por importación ({batch.file_name})"[:500],
                        )
                        EventService.log(
                            actor=acc.user,
                            action="collaborator.created",
                            supervisor=supervisor,
                            period=current.period,
                            target_type="collaborator",
                            target_id=collaborator.pk,
                            target_label=f"{collaborator.full_name} ({collaborator.cedula})",
                            new_value="ACTIVO",
                            source_file=batch.file_name,
                            context=context,
                            audit=False,
                        )
                        created.append(collaborator)
                    elif item["status"] == "EXISTING" and strategy == "update":
                        collaborator = Collaborator.objects.select_for_update().get(
                            pk=item["collaborator_id"]
                        )
                        for field, (old, new) in item["changes"].items():
                            setattr(collaborator, field, new)
                            EventService.log(
                                actor=acc.user,
                                action="collaborator.updated",
                                supervisor=supervisor,
                                period=current.period,
                                target_type="collaborator",
                                target_id=collaborator.pk,
                                target_label=f"{collaborator.full_name} ({collaborator.cedula})",
                                field=field,
                                old_value=old,
                                new_value=new,
                                source_file=batch.file_name,
                                context=context,
                                audit=False,
                            )
                        collaborator.save()
                        updated += 1
                    else:
                        unchanged += 1
            except IntegrityError:
                rejected.append(
                    {
                        "row": item["row"],
                        "column": "CEDULA",
                        "message": f"La cédula {item['cedula']} se registró en paralelo; "
                        "vuelva a importar.",
                    }
                )
        merged = {(e["row"], e["message"]): e for e in [*batch.errors, *rejected]}
        batch.status = ImportBatch.Status.CONFIRMED
        batch.strategy = strategy
        batch.errors = sorted(merged.values(), key=lambda e: e["row"])
        batch.summary = {
            **batch.summary,
            "created": len(created),
            "updated": updated,
            "skipped": unchanged,
            "rejected": len({e["row"] for e in rejected}),
            "applied": len(created) + updated,
        }
        batch.save(update_fields=["status", "strategy", "errors", "summary", "updated_at"])
        record_audit_event(
            actor=acc.user,
            action="novedades.collaborators.imported",
            target_type="import",
            target_id=batch.pk,
            module="novedades",
            new_values={**batch.summary, "file": batch.file_name},
            context=context,
        )
        EventService.log(
            actor=acc.user,
            action="import.confirmed",
            supervisor=supervisor,
            target_type="import",
            target_id=batch.pk,
            target_label=batch.file_name,
            new_value=f"{len(created)} creados, {updated} actualizados, {unchanged} sin cambios, "
            f"{batch.summary['rejected']} rechazados",
            source_file=batch.file_name,
            context=context,
            audit=False,
        )
        if created:
            # Alerta azul: un resumen por importación, solo con lo efectivamente
            # creado (las actualizaciones y los rechazos no son "nuevo ingreso").
            names = ", ".join(f"{c.full_name} ({c.cedula})" for c in created[:15])
            more = f" y {len(created) - 15} más" if len(created) > 15 else ""
            NotificationService.notify_admins(
                f"Nuevo(s) colaborador(es) registrado(s) en la base: {len(created)} — {names}{more}. "
                f"Equipo de {user_display(supervisor)}. Registrado por {user_display(acc.user)} el "
                f"{business_stamp()} (origen: importación «{batch.file_name}»; detalle en Historial "
                "de registros).",
                exclude_user=acc.user,
                category=ALERT_NEW_COLLABORATOR,
                dedup_key=f"novedades:collaborator-import:{batch.pk}",
            )
        if updated:
            NotificationService.notify_admins(
                f"Base de colaboradores: {user_display(acc.user)} actualizó {updated} colaborador(es) "
                f"del equipo de {user_display(supervisor)} por importación «{batch.file_name}».",
                exclude_user=acc.user,
                dedup_key=f"novedades:collaborator-import-upd:{batch.pk}",
            )
        return batch
